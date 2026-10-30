"""M5c: jumlah manpower per role dari Daily MP (--json) ke sheet ME di Subcontractor MP.xlsx.

Aturan: hanya MENAMBAH. Sel yang sudah berisi (isian manual) tidak pernah diganti.
  - dry run adalah default; menulis butuh apply=True / --apply
  - sel berisi angka berbeda -> berhenti dan lapor (skip_conflicts=True: lewati sel itu, tetap tidak menimpa)
  - satu jalur tulis: xlsx_sheet.fill_empty_cells; hasil ditulis ke FILE BARU, sumber tidak diubah
  - backup sumber ke data/backup/<nama>_<YYYYMMDD_HHMM>.xlsx sebelum menulis
  - verifikasi sesudah menulis: semua isian lama sama, jumlah isian naik tepat sebanyak yang ditulis
"""
import argparse
import json
import shutil
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

from ..common import config as cfgmod
from . import xlsx_sheet

ABSENT = {None, ""}


class VerifyError(Exception):
    pass


@dataclass
class Plan:
    date: str
    cells: dict = field(default_factory=dict)       # sel kosong yang akan diisi: ref -> angka
    same: list = field(default_factory=list)        # sel sudah berisi angka yang sama
    conflicts: list = field(default_factory=list)   # (ref, SC, role, di sheet, dari Daily MP)
    notes: list = field(default_factory=list)
    mapped: list = field(default_factory=list)      # role Daily MP yang dipetakan ke baris lain di sheet


def _col_letter(n):
    s = ""
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def date_col(cfg_ms, d):
    base = date.fromisoformat(cfg_ms["base_date"])
    return cfg_ms["base_col"] + (d - base).days


def _blocks(ws, cfg_ms):
    """{nama blok: (baris awal, baris akhir)} dari kolom B."""
    marks = [(r, ws.cell(r, 2).value) for r in range(cfg_ms["first_row"], cfg_ms["last_row"] + 1)
             if ws.cell(r, 2).value not in ABSENT and ws.cell(r, 2).value not in ("Indirect", "Direct")]
    out = {}
    for i, (r, name) in enumerate(marks):
        out[str(name).strip()] = (r, (marks[i + 1][0] if i + 1 < len(marks) else cfg_ms["last_row"] + 1) - 1)
    return out


_NEW_ROLE = " (role baru)"


def map_role(posisi, ms):
    """Role Daily MP -> baris di sheet. Role baru ('X (role baru)') mengikuti ms['role_map'][X], bila tidak ada
    masuk ms['new_role_default']; selain itu hanya alias ejaan."""
    if posisi.endswith(_NEW_ROLE):
        base = posisi[: -len(_NEW_ROLE)].strip()
        return ms.get("role_map", {}).get(base, ms.get("new_role_default", posisi))
    return ms.get("role_alias", {}).get(posisi, posisi)


def make_plan(xlsx, mp, cfg):
    """Bandingkan JSON Daily MP dengan sheet ME. Tidak menulis apa pun."""
    import openpyxl
    ms = cfg["me_sheet"]
    d = date.fromisoformat(mp["date"])
    plan = Plan(mp["date"])
    col = date_col(ms, d)
    hdr = openpyxl.load_workbook(xlsx, data_only=True)[ms["sheet"]].cell(ms["date_row"], col).value
    if hdr is None or hdr.date() != d:
        raise ValueError(f"kolom {_col_letter(col)} di sheet {ms['sheet']} bertanggal {hdr!r}, bukan {d}; tidak ditulis")
    ws = openpyxl.load_workbook(xlsx)[ms["sheet"]]
    blocks = _blocks(ws, ms)
    alias = ms.get("role_alias", {})
    for sc, name in ms["blocks"].items():
        if name not in blocks:
            plan.notes.append(f"{sc}: blok '{name}' tidak ada di kolom B sheet {ms['sheet']}")
            continue
        r0, r1 = blocks[name]
        rowof = {str(ws.cell(r, 4).value).strip(): r for r in range(r0, r1 + 1) if ws.cell(r, 4).value}
        agg = {}      # label di sheet -> jumlah; beberapa role Daily MP bisa menuju satu baris
        for row in mp["rows"]:
            v = row["values"].get(sc)
            if not isinstance(v, int) or v <= 0 or "tanpa rincian" in row["posisi"]:
                continue
            label = map_role(row["posisi"], ms)
            if label != row["posisi"] and label not in alias.values():
                plan.mapped.append(f"{sc}: '{row['posisi']}' ({v}) -> {label}")
            agg[label] = agg.get(label, 0) + v
        for label, v in agg.items():
            r = rowof.get(label)
            if r is None:
                plan.notes.append(f"{sc}: role '{label}' ({v}) tidak punya baris di blok {name}, tidak ditulis")
                continue
            ref, cur = f"{_col_letter(col)}{r}", ws.cell(r, col).value
            if cur in ABSENT:
                plan.cells[ref] = v
            elif cur == v:
                plan.same.append(ref)
            else:
                plan.conflicts.append((ref, sc, label, cur, v))
    for sc in mp["subs"]:
        if sc not in ms["blocks"] and mp["totals"].get(sc):
            plan.notes.append(f"{sc}: {mp['totals'][sc]} orang tidak ditulis (belum ada aturan blok untuk SC ini)")
    return plan


def report(plan):
    out = [f"Dry run ME {plan.date}: {len(plan.cells)} sel akan diisi, {len(plan.same)} sudah sama, "
           f"{len(plan.conflicts)} berbeda dari isian Anda (TIDAK diganti)"]
    out += [f"  isi  {ref} = {v}" for ref, v in plan.cells.items()]
    out += [f"  BEDA {ref} {sc} {role}: sheet={cur}, Daily MP={new}" for ref, sc, role, cur, new in plan.conflicts]
    out += [f"  dipetakan: {m}" for m in plan.mapped]
    out += [f"  catatan: {n}" for n in plan.notes]
    return "\n".join(out)


def _filled(xlsx, sheet):
    import openpyxl
    ws = openpyxl.load_workbook(xlsx)[sheet]
    return {c.coordinate: c.value for row in ws.iter_rows() for c in row if c.value is not None}


def apply_plan(xlsx, plan, cfg, dst=None, backup_dir=None, skip_conflicts=False, now=None):
    """Tulis plan ke FILE BARU (dst). Kembalikan path dst. Berhenti bila ada konflik (kecuali skip_conflicts)."""
    if plan.conflicts and not skip_conflicts:
        raise xlsx_sheet.OverwriteError(f"{len(plan.conflicts)} sel berbeda dari isian yang ada; berhenti. "
                                        "Periksa dry run, atau skip_conflicts untuk hanya mengisi sel kosong.")
    if not plan.cells:
        return None
    xlsx, sheet = Path(xlsx), cfg["me_sheet"]["sheet"]
    stamp = (now or datetime.now()).strftime("%Y%m%d_%H%M")
    backup_dir = Path(backup_dir or cfgmod.data_dir() / "backup")
    backup_dir.mkdir(parents=True, exist_ok=True)
    backup = backup_dir / f"{xlsx.stem}_{stamp}{xlsx.suffix}"
    shutil.copyfile(xlsx, backup)
    dst = Path(dst or xlsx.with_name(f"{xlsx.stem} - ME {plan.date}{xlsx.suffix}"))
    before = _filled(xlsx, sheet)
    xlsx_sheet.fill_empty_cells(xlsx, dst, sheet, plan.cells)
    after = _filled(dst, sheet)
    lost = [k for k, v in before.items() if after.get(k) != v]
    if lost:
        raise VerifyError(f"{len(lost)} isian lama berubah atau hilang: {', '.join(lost[:10])}")
    if len(after) != len(before) + len(plan.cells):
        raise VerifyError(f"jumlah isian {len(before)} -> {len(after)}, seharusnya +{len(plan.cells)}")
    return dst


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--xlsx", required=True)
    ap.add_argument("--json", required=True, help="keluaran dailymp.py --json")
    ap.add_argument("--apply", action="store_true", help="tulis (tanpa ini hanya dry run)")
    ap.add_argument("--skip-conflicts", action="store_true")
    ap.add_argument("--out")
    a = ap.parse_args(argv)
    cfg = cfgmod.load()
    plan = make_plan(a.xlsx, json.load(open(a.json, encoding="utf-8")), cfg)
    print(report(plan))
    if a.apply:
        dst = apply_plan(a.xlsx, plan, cfg, dst=a.out, skip_conflicts=a.skip_conflicts)
        print(f"ditulis: {dst}" if dst else "tidak ada sel baru untuk ditulis")


if __name__ == "__main__":
    main()

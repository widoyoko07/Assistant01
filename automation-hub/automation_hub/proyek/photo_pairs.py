"""M4b: butir Work Description terpilih ("x" di kolom Pakai, sheet WD Kandidat) -> foto, per (tanggal, SC).

Butir tidak wajib punya foto. Tanggal yang sudah punya cukup foto (folder Drive foto kumpulan pengguna, atau baris 'foto'
di sheet Work Description) tidak perlu dikumpulkan lagi; gaps() menunjukkan tanggal yang masih kurang.

Pasangan dibuat menurut urutan: butir ke-n ke foto ke-n pada laporan hari itu. Ini usulan awal; tabel hasilnya
dikoreksi pengguna. Butir tanpa foto dan foto sisa dilaporkan, tidak ditebak.
"""
import csv
from pathlib import Path

from ..common import config as cfgmod
from ..common.runlog import RunLog

DRIVE_FULL = 6    # foto per tanggal di folder Drive yang dianggap lengkap
SHEET_FULL = 4    # baris 'foto' di sheet Work Description; di bawah ini berarti foto belum lengkap
HEADER = ["Tanggal", "SC", "No", "Butir Final", "Foto", "Catatan"]


def selected_items(xlsx, sheet="WD Kandidat", fixes=None):
    """[(tanggal, SC, butir)] untuk baris ber-"x" di kolom Pakai. Butir = kolom 'Butir Final' bila terisi, selain itu 'Butir'.
    fixes {salah: benar} memperbaiki salah ketik yang sudah dikonfirmasi pengguna."""
    import openpyxl
    wb = openpyxl.load_workbook(xlsx, read_only=True)
    rows = wb[sheet].iter_rows(values_only=True)
    head = list(next(rows))
    col = {h: i for i, h in enumerate(head) if h}
    out = []
    for r in rows:
        if str(r[col["Pakai"]] or "").strip().lower() != "x":
            continue
        final = r[col["Butir Final"]] if "Butir Final" in col else None
        text = str(final or r[col["Butir"]]).strip()
        for bad, good in (fixes or {}).items():
            text = text.replace(bad, good)
        out.append((r[col["Tanggal"]], r[col["SC"]], text))
    return out


def drive_counts(names):
    """{YYYY-MM-DD: jumlah} dari nama berkas foto di folder Drive ('YYMMDD Judul.jpg')."""
    out = {}
    for n in names:
        if len(n) > 6 and n[:6].isdigit():
            d = f"20{n[:2]}-{n[2:4]}-{n[4:6]}"
            out[d] = out.get(d, 0) + 1
    return out


def sheet_counts(xlsx, sheet="Work Description"):
    """{YYYY-MM-DD: nilai baris 'foto'} dari sheet Work Description (baris 2 = tanggal, baris dengan label 'foto' di kolom A)."""
    import datetime
    import openpyxl
    ws = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)[sheet]
    dates, foto = None, None
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=6, values_only=True)):
        if i == 1:
            dates = row
        if str(row[0] or "").strip().lower() == "foto":
            foto = row
    out = {}
    for d, v in zip(dates or (), foto or ()):
        if isinstance(d, datetime.datetime):
            out[d.date().isoformat()] = v if isinstance(v, (int, float)) else 0
    return out


def gaps(dates, drive, sheet):
    """[(tanggal, foto di Drive, foto di sheet)] untuk tanggal yang masih perlu dikumpulkan: Drive < DRIVE_FULL dan sheet < SHEET_FULL."""
    return [(d, drive.get(d, 0), sheet.get(d, 0)) for d in dates
            if drive.get(d, 0) < DRIVE_FULL and sheet.get(d, 0) < SHEET_FULL]


def manifest_photos(root=None):
    """{(tanggal, SC): [foto]} dari data/04_photos/manifest.csv, urutan laporan."""
    base = cfgmod.data_dir(root=root) / "04_photos"
    found = {}
    path = base / "manifest.csv"
    if not path.is_file():
        return found
    with open(path, encoding="utf-8-sig", newline="") as f:
        for r in csv.DictReader(f):
            if r["file"]:
                found.setdefault((r["date"], r["sc"]), []).append(str(base / r["date"].replace("-", "") / r["sc"] / r["file"]))
    return found


def pair(items, photos):
    """items: [(tanggal, SC, butir)], photos: {(tanggal, SC): [path]} -> baris tabel (lihat HEADER)."""
    by_key = {}
    for d, sc, text in items:
        by_key.setdefault((d, sc), []).append(text)
    rows = []
    for (d, sc), texts in sorted(by_key.items()):
        pics = photos.get((d, sc), [])
        for n, text in enumerate(texts):
            if n < len(pics):
                rows.append([d, sc, n + 1, text, pics[n], ""])
            else:
                rows.append([d, sc, n + 1, text, "", "tanpa foto" if not pics else "tanpa foto"])
        for extra in pics[len(texts):]:
            rows.append([d, sc, "", "", extra, "foto sisa (belum dipasangkan)"])
    return rows


def write_csv(xlsx, root=None, log=None):
    log = log or RunLog("photo_pairs", root=root)
    rows = pair(selected_items(xlsx, fixes=cfgmod.load().get("butir_fix")), manifest_photos(root))
    out_dir = cfgmod.data_dir(root=root) / "05_pairs"
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / "pasangan_foto_wd.csv"
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(HEADER)
        w.writerows(rows)
    missing = sum(1 for r in rows if r[3] and not r[4])
    if missing:
        log.warn(f"{missing} butir terpilih tanpa foto")
    log.done(f"{path} ({len(rows)} baris)")
    return path


if __name__ == "__main__":
    import sys
    write_csv(Path(sys.argv[1]))

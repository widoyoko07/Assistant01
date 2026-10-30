import tempfile
import unittest
from pathlib import Path

import openpyxl

from automation_hub.proyek import xlsx_sheet as X


class AddSheetTest(unittest.TestCase):
    def make(self, d):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Work Description"
        ws["A1"] = "=SUM(B1:B2)"
        ws["B1"], ws["B2"] = 2, 3
        wb.create_sheet("ME")["A1"] = "x"
        p = Path(d) / "src.xlsx"
        wb.save(p)
        return p

    def test_adds_sheet_and_keeps_others(self):
        with tempfile.TemporaryDirectory() as d:
            src, dst = self.make(d), Path(d) / "dst.xlsx"
            X.add_sheet(src, dst, "WD Kandidat", [["Tanggal", "Butir"], ["2026-10-24", "- A & B <c>"], ["2026-10-24", 5]])
            wb = openpyxl.load_workbook(dst)
            self.assertEqual(wb.sheetnames, ["Work Description", "ME", "WD Kandidat"])
            self.assertEqual(wb["Work Description"]["A1"].value, "=SUM(B1:B2)")
            ws = wb["WD Kandidat"]
            self.assertEqual([c.value for c in ws[2]], ["2026-10-24", "- A & B <c>"])
            self.assertEqual(ws["B3"].value, 5)
            self.assertEqual(ws.freeze_panes, "A2")

    def test_existing_name_refused_and_source_untouched(self):
        with tempfile.TemporaryDirectory() as d:
            src = self.make(d)
            before = src.read_bytes()
            with self.assertRaises(ValueError):
                X.add_sheet(src, Path(d) / "dst.xlsx", "ME", [["a"]])
            with self.assertRaises(ValueError):
                X.add_sheet(src, Path(d) / "dst.xlsx", "a/b", [["a"]])
            self.assertEqual(src.read_bytes(), before)


if __name__ == "__main__":
    unittest.main()

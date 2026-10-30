import tempfile
import unittest
from datetime import date, datetime, timedelta
from pathlib import Path

import openpyxl

from automation_hub.common import config
from automation_hub.proyek import me_sheet as MS
from automation_hub.proyek import xlsx_sheet as X

CFG = config.load()
BASE = date.fromisoformat(CFG["me_sheet"]["base_date"])
COL = CFG["me_sheet"]["base_col"]


def make(d, filled=None):
    """Sheet ME sintetis: tanggal di baris 4, blok TPE di baris 313 dengan 3 role."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ME"
    for i in range(0, 5):
        ws.cell(4, COL + i).value = datetime.combine(BASE, datetime.min.time()) + timedelta(days=i)
    ws.cell(313, 2).value = "TPE"
    for k, role in enumerate(["Supervisor", "Skilled Workers", "Commoun Labor"]):
        ws.cell(314 + k, 4).value = role
    ws.cell(320, 2).value = "WME"
    ws.cell(321, 4).value = "Supervisor"
    for (r, c), v in (filled or {}).items():
        ws.cell(r, c).value = v
    p = Path(d) / "smp.xlsx"
    wb.save(p)
    return p


def mp(day, **vals):
    rows = [{"type": "", "posisi": k.replace("_", " "), "values": {"TPE": v}} for k, v in vals.items()]
    return {"date": day.isoformat(), "subs": ["TPE"], "rows": rows, "totals": {"TPE": sum(vals.values())}}


class FillEmptyCellsTest(unittest.TestCase):
    def test_refuses_nonempty_and_writes_nothing(self):
        with tempfile.TemporaryDirectory() as t:
            src = make(t, {(314, COL + 1): 5})
            dst = Path(t) / "o.xlsx"
            with self.assertRaises(X.OverwriteError):
                X.fill_empty_cells(src, dst, "ME", {"F314": 9, "F315": 1})
            self.assertFalse(dst.exists())

    def test_src_and_dst_must_differ_and_values_numeric(self):
        with tempfile.TemporaryDirectory() as t:
            src = make(t)
            with self.assertRaises(ValueError):
                X.fill_empty_cells(src, src, "ME", {"F314": 1})
            with self.assertRaises(ValueError):
                X.fill_empty_cells(src, Path(t) / "o.xlsx", "ME", {"F314": "x"})

    def test_fills_new_cells_keeps_rest(self):
        with tempfile.TemporaryDirectory() as t:
            src = make(t)
            dst = Path(t) / "o.xlsx"
            X.fill_empty_cells(src, dst, "ME", {"F314": 3, "AZ500": 2})
            ws = openpyxl.load_workbook(dst)["ME"]
            self.assertEqual((ws["F314"].value, ws["AZ500"].value, ws["D315"].value), (3, 2, "Skilled Workers"))


class PlanApplyTest(unittest.TestCase):
    D = BASE + timedelta(days=1)   # kolom F

    def test_plan_fills_empty_marks_same_and_conflict_alias_applied(self):
        with tempfile.TemporaryDirectory() as t:
            src = make(t, {(314, COL + 1): 4, (315, COL + 1): 7})
            plan = MS.make_plan(src, mp(self.D, Supervisor=4, Skilled_Workers=9, Common_Labor=2), CFG)
            self.assertEqual(plan.cells, {"F316": 2})            # 'Common Labor' -> 'Commoun Labor'
            self.assertEqual(plan.same, ["F314"])
            self.assertEqual([c[:2] for c in plan.conflicts], [("F315", "TPE")])

    def test_apply_stops_on_conflict_unless_skipped_and_never_overwrites(self):
        with tempfile.TemporaryDirectory() as t:
            src = make(t, {(315, COL + 1): 7})
            plan = MS.make_plan(src, mp(self.D, Supervisor=4, Skilled_Workers=9), CFG)
            with self.assertRaises(X.OverwriteError):
                MS.apply_plan(src, plan, CFG, backup_dir=Path(t) / "bk")
            dst = MS.apply_plan(src, plan, CFG, dst=Path(t) / "o.xlsx", backup_dir=Path(t) / "bk", skip_conflicts=True)
            ws = openpyxl.load_workbook(dst)["ME"]
            self.assertEqual((ws["F314"].value, ws["F315"].value), (4, 7))
            self.assertEqual(len(list((Path(t) / "bk").glob("smp_*.xlsx"))), 1)   # backup ada

    def test_wrong_date_header_refused(self):
        with tempfile.TemporaryDirectory() as t:
            src = make(t)
            with self.assertRaises(ValueError):
                MS.make_plan(src, mp(BASE + timedelta(days=30), Supervisor=1), CFG)

    def test_dry_run_default_writes_nothing(self):
        with tempfile.TemporaryDirectory() as t:
            src = make(t)
            before = src.read_bytes()
            plan = MS.make_plan(src, mp(self.D, Supervisor=4), CFG)
            self.assertIn("Dry run", MS.report(plan))
            self.assertEqual(src.read_bytes(), before)


class SingleWritePathTest(unittest.TestCase):
    def test_me_sheet_has_no_direct_writers(self):
        src = (Path(MS.__file__)).read_text(encoding="utf-8")
        for banned in (".save(", "ZipFile(", "write_text(", "write_bytes(", "open(path, \"w\"", "writestr("):
            self.assertNotIn(banned, src, f"me_sheet.py tidak boleh menulis langsung ({banned}); pakai xlsx_sheet.fill_empty_cells")


if __name__ == "__main__":
    unittest.main()

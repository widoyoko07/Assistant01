#!/usr/bin/env python3
"""Conduit per UKURAN dan JENIS (Power/Signal/Control) dari PDF shop drawing vektor.

Label conduit ("42mmØ / SIGNAL / EL-12838mm") digambar sebagai garis vektor (bukan teks PDF), jadi dibaca dengan OCR
pada render halaman. Modul ini memakai hasil batch tray/conduit yang sudah ada (pages.jsonl + plan.json di --plan-dir).

  python conduit.py plan    PDF --plan-dir DIR_BATCH --out DIR [--budget 240] [--dpi 300] [--pages 28,490]
  python conduit.py execute PDF --plan-dir DIR_BATCH --out DIR --confirm

plan    : OCR halaman berconduit (checkpoint conduit_labels.jsonl, ulangi perintah sama untuk melanjutkan), lalu
          tabel dry-run per set: jumlah label per ukuran x jenis dan meter ESTIMASI.
execute : tulis Rekap_Conduit_per_Area.xlsx (format sama dengan rekap tray: baris = area, kolom = ukuran x jenis).

Metode (jujur): jumlah label = hasil baca OCR (bisa meleset bila teks samar); meter per kategori = panjang conduit
hasil ukur garis hijau di halaman itu, DIALOKASIKAN PROPORSIONAL terhadap jumlah label di halaman yang sama.
Itu estimasi, bukan penelusuran leader per label. Conduit vertikal dari section = batas atas (belum dedup)."""
import argparse, collections, difflib, json, os, re, sys
import numpy as np

HERE = os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0, HERE)
TYPES = ("POWER", "SIGNAL", "CONTROL")
KNOWN_SIZES = (16, 22, 28, 36, 42, 54, 70, 82, 92, 104)   # ukuran umum di gambar proyek; selain ini ditandai
SIZE_RE = re.compile(r"^([0-9OoIl]{2,3})\s*mm", re.I)
DIGITS = str.maketrans("OoIiLl", "001111")
LOW_CONF = 30
KIND_LABEL = "Jumlah label"; KIND_H = "Horizontal m (estimasi, plan)"; KIND_V = "Vertikal m (batas atas, section)"


class ConduitStateError(RuntimeError):
    pass


# ------------------------------------------------------------ OCR
def ocr_words(page, dpi=300, xmax_frac=0.9):
    """Render halaman, tebalkan garis tipis, OCR -> [{t,x,y,w,h,c}] (piksel). Label vektor abu-abu tipis butuh threshold + erode."""
    import cv2, pytesseract
    pm = page.get_pixmap(dpi=dpi, colorspace=__import__("pymupdf").csGRAY)
    im = np.frombuffer(pm.samples, np.uint8).reshape(pm.height, pm.width)
    _, b = cv2.threshold(im, 225, 255, cv2.THRESH_BINARY)
    b = cv2.erode(b, np.ones((2, 2), np.uint8))[:, :int(pm.width * xmax_frac)]
    try:
        d = pytesseract.image_to_data(b, lang="eng", config="--psm 11", output_type=pytesseract.Output.DICT)
    except pytesseract.pytesseract.TesseractNotFoundError as exc:
        raise RuntimeError("Tesseract belum terpasang atau tidak ada di PATH.") from exc
    return [dict(t=d["text"][i].strip(), x=d["left"][i], y=d["top"][i], w=d["width"][i], h=d["height"][i], c=float(d["conf"][i]))
            for i in range(len(d["text"])) if d["text"][i].strip()]


# ------------------------------------------------------------ parse label
def type_of(tok):
    s = re.sub(r"[^A-Z]", "", tok.upper())
    if len(s) < 4: return None
    best = max(TYPES, key=lambda t: difflib.SequenceMatcher(None, s, t).ratio())
    return best if difflib.SequenceMatcher(None, s, best).ratio() >= 0.72 else None


def parse_labels(words):
    """Pasangkan token ukuran ('42mmØ') dengan token jenis tepat di bawahnya (rata kiri). Return (labels, unpaired_sizes).
    Token yang diawali EL (elevasi) atau berdigit >3 diabaikan. Ukuran tanpa jenis dilaporkan, tidak dibuang diam-diam."""
    sizes, unpaired = [], []
    for w in words:
        if w["t"].upper().lstrip("_").startswith("EL"): continue
        m = SIZE_RE.match(w["t"])
        if m:
            normalized = m.group(1).translate(DIGITS)
            if normalized.isdigit(): sizes.append((int(normalized), w))
            else: unpaired.append({"size": m.group(1), "x": w["x"], "y": w["y"], "conf": w["c"]})
    types = [(type_of(w["t"]), w) for w in words]; types = [(t, w) for t, w in types if t]
    used, labels = set(), []
    for size, w in sorted(sizes, key=lambda s: (s[1]["y"], s[1]["x"])):
        best = None
        for j, (t, tw) in enumerate(types):
            dy = tw["y"] - w["y"]
            if j in used or not (0.3 * w["h"] <= dy <= 3.2 * w["h"] or 15 <= dy <= 110): continue
            if abs(tw["x"] - w["x"]) > max(60, 2.0 * w["h"]): continue
            if best is None or dy < best[0]: best = (dy, j, t)
        if best is None: unpaired.append({"size": size, "x": w["x"], "y": w["y"], "conf": w["c"]}); continue
        used.add(best[1])
        labels.append({"size": size, "type": best[2], "x": w["x"], "y": w["y"], "conf": w["c"],
                       "known": size in KNOWN_SIZES, "low_conf": w["c"] < LOW_CONF})
    return labels, unpaired


# ------------------------------------------------------------ alokasi + agregasi
def allocate(counts, length_m):
    """Bagi panjang halaman ke kategori (size,type) proporsional terhadap jumlah label. Tanpa label -> {} (tidak teridentifikasi)."""
    tot = sum(counts.values())
    return {k: length_m * v / tot for k, v in counts.items()} if tot and length_m else {}


def page_lengths(rec):
    """(kind, meter) dari hasil batch: plan -> horizontal; section -> vertikal (batas atas)."""
    if not rec: return None, 0.0
    if rec["role"] == "plan": return "h", float(rec.get("conduit_h_m", 0) or 0)
    if rec["role"] == "section": return "v", float(rec.get("conduit_v_m", 0) or 0)
    return None, 0.0


def aggregate(plan_dir, labels_by_page):
    """-> list per set: {set, pages, name, include, count{cat:n}, h{cat:m}, v{cat:m}, h_none, v_none, flags}."""
    with open(os.path.join(plan_dir, "plan.json"), encoding="utf-8") as fh: plan = json.load(fh)
    latest = {}
    with open(os.path.join(plan_dir, "pages.jsonl"), encoding="utf-8") as fh:
        for line in fh:
            it = json.loads(line); latest[it["row"]["sheet"]] = it
    out = []
    for s in plan["sets"]:
        c = collections.Counter(); h = collections.defaultdict(float); v = collections.defaultdict(float)
        none = {"h": 0.0, "v": 0.0}; flags = []; low = 0; unk = collections.Counter(); unp = 0; skipped = []
        for pg in range(s["pages"][0], s["pages"][1] + 1):
            it = latest.get(pg)
            if not it or not it["row"]["items"].get("conduit"): continue
            kind, length = page_lengths(it["rec"])
            lab = labels_by_page.get(pg)
            if lab is None: skipped.append(pg); continue
            cnt = collections.Counter((l["size"], l["type"]) for l in lab["labels"])
            c.update(cnt); low += sum(1 for l in lab["labels"] if l["low_conf"]); unp += len(lab["unpaired"])
            unk.update(l["size"] for l in lab["labels"] if not l["known"])
            if kind:
                share = allocate(cnt, length)
                tgt = h if kind == "h" else v
                for k, m in share.items(): tgt[k] += m
                if length and not share: none[kind] += length
        if skipped: flags.append(f"{len(skipped)} halaman berconduit belum di-OCR/tidak terukur: hal {skipped[:10]}")
        if low: flags.append(f"{low} label confidence OCR <{LOW_CONF}: cek visual")
        if unp: flags.append(f"{unp} token ukuran tanpa jenis di bawahnya (tidak dihitung)")
        if unk: flags.append("ukuran di luar daftar umum (mungkin salah baca OCR): " + ", ".join(f"{k}mm x{n}" for k, n in sorted(unk.items())))
        if none["h"] or none["v"]: flags.append(f"conduit tanpa label terbaca: horizontal {none['h']:.1f} m, vertikal {none['v']:.1f} m")
        out.append({"set": s["set"], "pages": s["pages"], "name": s["name"], "include": s["include"], "count": dict(c),
                    "h": dict(h), "v": dict(v), "h_none": round(none["h"], 1), "v_none": round(none["v"], 1), "flags": flags})
    return out


def cat_str(k): return f"{k[0]}|{k[1]}"
def cat_parse(s): a, b = s.split("|"); return int(a), b


# ------------------------------------------------------------ OCR runner (checkpoint)
def conduit_pages(plan_dir, only=None):
    pages = []
    with open(os.path.join(plan_dir, "pages.jsonl"), encoding="utf-8") as fh:
        for line in fh:
            it = json.loads(line)
            if it["row"]["items"].get("conduit") and it["row"]["role"] in ("plan", "section") and (not only or it["row"]["sheet"] in only):
                pages.append(it["row"]["sheet"])
    return sorted(set(pages))


def run_ocr(pdf, plan_dir, out, dpi=300, budget=240, only=None):
    import time, pymupdf as fitz, takeoff as T
    os.makedirs(out, exist_ok=True)
    path = os.path.join(out, "conduit_labels.jsonl"); run = os.path.join(out, "conduit_run.json")
    manifest = T.source_manifest(pdf, T.load_config(), mode="conduit-ocr", dpi=dpi, sizes=list(KNOWN_SIZES), threshold=225, erode=2)
    if os.path.exists(run):
        with open(run, encoding="utf-8") as fh: prev = json.load(fh)
        if prev != manifest:
            raise ConduitStateError("Checkpoint OCR berasal dari file/opsi berbeda. Gunakan folder output baru.")
    elif os.path.exists(path): raise ConduitStateError("Checkpoint lama tanpa conduit_run.json. Gunakan folder output baru.")
    else:
        with open(run, "w", encoding="utf-8") as fh: json.dump(manifest, fh, indent=1)
    done = {}
    if os.path.exists(path):
        with open(path, encoding="utf-8") as fh:
            for n, ln in enumerate(fh, 1):
                try: r = json.loads(ln); done[r["sheet"]] = r
                except (json.JSONDecodeError, KeyError) as e: raise ConduitStateError(f"Checkpoint rusak baris {n}; gunakan folder output baru.") from e
    todo = [p for p in conduit_pages(plan_dir, only) if p not in done]; doc = fitz.open(pdf); t0 = time.time()
    with open(path, "a", encoding="utf-8") as fh:
        for pg in todo:
            if time.time() - t0 > budget: break
            lab, unp = parse_labels(ocr_words(doc[pg - 1], dpi))
            r = {"sheet": pg, "labels": lab, "unpaired": unp}; fh.write(json.dumps(r, ensure_ascii=False) + "\n"); fh.flush(); done[pg] = r
    left = [p for p in conduit_pages(plan_dir, only) if p not in done]
    print(f"OCR: {len(conduit_pages(plan_dir, only)) - len(left)}/{len(conduit_pages(plan_dir, only))} halaman berconduit.")
    return done, manifest, not left


# ------------------------------------------------------------ tampilan + Excel
def categories(rows):
    sizes = sorted({k[0] for r in rows for d in (r["count"], r["h"], r["v"]) for k in d})
    return [(s, t) for s in sizes for t in TYPES]


def print_plan(rows):
    """Dry-run ringkas: total per ukuran x jenis (set 'ya'), lalu ringkasan per set. Hanya baris/kolom yang berisi."""
    inc = [r for r in rows if r["include"]]
    cnt = collections.Counter(); mtr = collections.defaultdict(float)
    for r in inc:
        for k, v in r["count"].items(): cnt[k] += v
        for d in (r["h"], r["v"]):
            for k, v in d.items(): mtr[k] += v
    print("## DRY-RUN conduit per ukuran x jenis (belum ada Excel)\n")
    print("Total set 'ya': jumlah label / meter estimasi (horizontal plan + vertikal section batas atas)\n")
    print("| Ukuran | Power | Signal | Control | Total |\n|---|---|---|---|---|")
    for size in sorted({k[0] for k in list(cnt) + list(mtr)}):
        cells = [f"{cnt.get((size, t), 0)} / {mtr.get((size, t), 0):.0f} m" for t in TYPES]
        print(f"| {size} mm | " + " | ".join(cells) + f" | {sum(cnt.get((size, t), 0) for t in TYPES)} / {sum(mtr.get((size, t), 0) for t in TYPES):.0f} m |")
    print(f"\nTotal label {sum(cnt.values())} | meter teralokasi {sum(mtr.values()):.0f} m | tanpa label: "
          f"horizontal {sum(r['h_none'] for r in inc):.0f} m, vertikal {sum(r['v_none'] for r in inc):.0f} m")
    print("\n| Set | Hal | Nama | Hitung | Label | Meter H | Meter V (batas atas) | Tanpa label H/V |\n|---|---|---|---|---|---|---|---|")
    for r in rows:
        print(f"| {r['set']} | {r['pages'][0]}-{r['pages'][1]} | {r['name'][:30]} | {'ya' if r['include'] else 'TIDAK'} | {sum(r['count'].values())} | "
              f"{sum(r['h'].values()):.0f} | {sum(r['v'].values()):.0f} | {r['h_none']}/{r['v_none']} |")
    print("\nPeringatan:")
    for r in rows:
        for f in r["flags"]: print(f" - set {r['set']} ({r['name'][:28]}): {f}")


def write_excel(out, rows, area_label):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter as CL
    cats = categories(rows); n = len(cats)
    for r in rows: r["area"] = area_label(r["name"])
    areas = []
    for r in rows:
        if r["area"] not in areas: areas.append(r["area"])
    F = Font(name="Arial"); FB = Font(name="Arial", bold=True); HD = PatternFill("solid", fgColor="305496")
    thin = Side(style="thin", color="999999"); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
    wb = Workbook(); wr = wb.active; wr.title = "Rekap Conduit per Area"; wd = wb.create_sheet("Detail Set")
    wd.append(["Set", "Area", "Hal awal", "Hal akhir", "Hitung (Y/N)", "Metrik"] + [f"{s}mm {t}" for s, t in cats] + ["Tanpa label (m)", "Catatan"])
    for c in wd[1]: c.font = Font(name="Arial", bold=True, color="FFFFFF"); c.fill = HD; c.alignment = Alignment(wrap_text=True, vertical="center")
    for r in rows:
        base = [r["set"], r["area"], r["pages"][0], r["pages"][1], "Y" if r["include"] else "N"]
        wd.append(base + [KIND_LABEL] + [r["count"].get(k, 0) for k in cats] + [0, "; ".join(r["flags"])[:700]])
        wd.append(base + [KIND_H] + [round(r["h"].get(k, 0), 2) for k in cats] + [r["h_none"], ""])
        wd.append(base + [KIND_V] + [round(r["v"].get(k, 0), 2) for k in cats] + [r["v_none"], ""])
    nd = wd.max_row; wd.freeze_panes = "G2"
    for i in range(1, wd.max_column + 1): wd.column_dimensions[CL(i)].width = 90 if i == wd.max_column else (32 if i in (2, 6) else 11)
    for row in wd.iter_rows(min_row=2):
        for c in row: c.font = F
    r0 = 1
    wr.cell(r0, 1, "Rekap Conduit per Area: ukuran x jenis").font = Font(name="Arial", bold=True, size=13); r0 = 3
    last_c = 1 + n + 2; blocks = [(KIND_LABEL, "Jumlah label (OCR)", "0"), (KIND_H, "Panjang horizontal (m, estimasi dari plan)", "0.0"),
                                  (KIND_V, "Panjang vertikal (m, batas atas dari section)", "0.0")]
    for kind, title, fmt in blocks:
        wr.cell(r0, 1, title).font = FB
        sizes = sorted({s for s, _ in cats}); hr = r0 + 1
        for j, s in enumerate(sizes):
            c0 = 2 + j * len(TYPES); wr.cell(hr, c0, f"{s} mm"); wr.merge_cells(start_row=hr, start_column=c0, end_row=hr, end_column=c0 + len(TYPES) - 1)
            for q, t in enumerate(TYPES): wr.cell(hr + 1, c0 + q, t.title())
        wr.cell(hr + 1, 1, "Area")
        extra = 2 + n
        if kind != KIND_LABEL: wr.cell(hr + 1, extra, "Tanpa label")
        wr.cell(hr + 1, extra + (0 if kind == KIND_LABEL else 1), "Total")
        endc = extra + (0 if kind == KIND_LABEL else 1)
        for rr in (hr, hr + 1):
            for cc in range(1, endc + 1):
                c = wr.cell(rr, cc); c.font = Font(name="Arial", bold=True, color="FFFFFF"); c.fill = HD
                c.alignment = Alignment(horizontal="center", vertical="center"); c.border = BOX
        for i, a in enumerate(areas):
            r = hr + 2 + i; wr.cell(r, 1, a)
            for j in range(n + (0 if kind == KIND_LABEL else 1)):
                col = CL(7 + j)
                wr.cell(r, 2 + j, f"=SUMIFS('Detail Set'!${col}$2:${col}${nd},'Detail Set'!$B$2:$B${nd},$A{r},'Detail Set'!$E$2:$E${nd},\"Y\",'Detail Set'!$F$2:$F${nd},\"{kind}\")").number_format = fmt
            wr.cell(r, endc, f"=SUM({CL(2)}{r}:{CL(endc - 1)}{r})").number_format = fmt
            for cc in range(1, endc + 1): wr.cell(r, cc).border = BOX; wr.cell(r, cc).font = F
        rt = hr + 2 + len(areas); wr.cell(rt, 1, "TOTAL")
        for cc in range(2, endc + 1): wr.cell(rt, cc, f"=SUM({CL(cc)}{hr + 2}:{CL(cc)}{rt - 1})").number_format = fmt
        for cc in range(1, endc + 1): wr.cell(rt, cc).font = FB; wr.cell(rt, cc).border = BOX
        r0 = rt + 3
    notes = ["Catatan:",
             "1. Jumlah label = hasil OCR label conduit (ukuran + jenis). Label bisa terlewat bila teks samar; baris yang ditandai ada di sheet Detail Set kolom Catatan.",
             "2. Meter per ukuran x jenis = ESTIMASI: panjang conduit hijau di halaman dibagi proporsional jumlah label di halaman itu (bukan penelusuran leader per label).",
             "3. 'Tanpa label' = panjang conduit pada halaman yang tidak punya label terbaca; belum diklasifikasi.",
             "4. Vertikal dari section = batas atas (conduit yang sama bisa tampil di beberapa section; label plan dan section bisa menghitung conduit yang sama dua kali).",
             "5. Set duplikat/revisi diatur di sheet Detail Set kolom Hitung (Y/N); rekap terhitung ulang otomatis."]
    for i, t in enumerate(notes): wr.cell(r0 + i, 1, t).font = FB if i == 0 else F
    wr.column_dimensions["A"].width = 30
    for cc in range(2, last_c + 2): wr.column_dimensions[CL(cc)].width = 11
    wr.freeze_panes = "B3"
    path = os.path.join(out, "Rekap_Conduit_per_Area.xlsx"); wb.save(path); print("Excel:", path); return path


# ------------------------------------------------------------ CLI
def load_labels(out):
    p = os.path.join(out, "conduit_labels.jsonl")
    if not os.path.exists(p): return {}
    with open(p, encoding="utf-8") as fh: return {json.loads(l)["sheet"]: json.loads(l) for l in fh}


def main():
    import takeoff as T, batch
    ap = argparse.ArgumentParser(); sp = ap.add_subparsers(dest="cmd", required=True)
    for n in ("plan", "execute"):
        s = sp.add_parser(n); s.add_argument("pdf"); s.add_argument("--plan-dir", required=True, help="folder hasil batch takeoff (pages.jsonl + plan.json)")
        s.add_argument("--out", required=True); s.add_argument("--dpi", type=int, default=300); s.add_argument("--budget", type=int, default=240)
        s.add_argument("--pages", help="batasi ke halaman tertentu, mis. 28,490 (uji)")
        if n == "execute": s.add_argument("--confirm", action="store_true")
    a = ap.parse_args(); only = {int(x) for x in a.pages.split(",")} if a.pages else None
    for f in ("pages.jsonl", "plan.json"):
        if not os.path.exists(os.path.join(a.plan_dir, f)): sys.exit(f"{f} tidak ada di {a.plan_dir}: jalankan takeoff.py plan (batch) dulu.")
    try:
        if a.cmd == "plan":
            labels, _, complete = run_ocr(a.pdf, a.plan_dir, a.out, a.dpi, a.budget, only)
            if not complete: print("BELUM selesai: jalankan perintah yang sama lagi untuk melanjutkan."); return
            print_plan(aggregate(a.plan_dir, labels)); return
        if not a.confirm: sys.exit("Ditolak: execute hanya setelah user menyetujui hasil dry-run (tambahkan --confirm).")
        run = os.path.join(a.out, "conduit_run.json")
        if not os.path.exists(run): sys.exit("conduit_run.json tidak ada: jalankan plan dulu.")
        exp = T.source_manifest(a.pdf, T.load_config(), mode="conduit-ocr", dpi=a.dpi, sizes=list(KNOWN_SIZES), threshold=225, erode=2)
        with open(run, encoding="utf-8") as fh: stored = json.load(fh)
        if stored != exp: sys.exit("Sumber atau opsi tidak cocok dengan dry-run. Jalankan plan ulang di folder output yang sama.")
        labels = load_labels(a.out); missing = [p for p in conduit_pages(a.plan_dir, only) if p not in labels]
        if missing: sys.exit(f"OCR belum lengkap ({len(missing)} halaman). Jalankan plan sampai selesai.")
        write_excel(a.out, aggregate(a.plan_dir, labels), batch.area_label)
    except ConduitStateError as e: sys.exit(str(e))


if __name__ == "__main__": main()

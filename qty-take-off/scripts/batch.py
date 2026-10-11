#!/usr/bin/env python3
"""Mode batch qty-take-off: PDF gabungan banyak sheet/set gambar (mis. 599 halaman). Checkpoint per halaman (pages.jsonl),
dapat dilanjutkan (jalankan perintah yang sama lagi), hasil dikelompokkan per SET gambar (berdasar halaman DRAWING INDEX),
duplikat versi ditandai agar tidak terhitung ganda. Tray/conduit saja."""
import collections, json, os, re, sys, time
import numpy as np
import takeoff as T

BIN = 50                                   # lebar tray dikelompokkan per 50 mm sebelum dipetakan ke lebar kandidat
KEYWORDS = ("TANKS", "BUILDING", "ROOM", "HEADWORKS", "FACILITY", "CHAMBER", "FILTER", "STATION", "SLUDGE", "CHLORINE", "MBR", "A2O")
SKIP_WORDS = re.compile(r"\b(PANELS?|AND|TRANSFORMERS?|CABLE|TRAY|CONDUIT|BUSDUCT|LAYOUT|FIELD|EQUIPMENTS?|INCLUDE|ROUTING|&)\b", re.I)


def hist_measure(m, close_pt, px, mm_pt, thick_pt=0.6):
    """Garis tengah + histogram lebar lokal -> {lebar_mm_bin: meter}. Lebar = (lebar_pt - tebal garis) * mm/pt."""
    import cv2
    from skimage.morphology import skeletonize
    from scipy import ndimage as ndi
    if not m.any(): return {}, m > 0
    k = max(1, round(close_pt * px))
    mc = cv2.morphologyEx(m, cv2.MORPH_CLOSE, cv2.getStructuringElement(cv2.MORPH_RECT, (k, k)))
    sk = skeletonize(mc > 0); ys, xs = np.nonzero(sk)
    if not len(xs): return {}, sk
    w_mm = np.clip((2 * ndi.distance_transform_edt(mc > 0)[ys, xs] / px - thick_pt) * mm_pt, 0, 1500)
    cls = np.zeros(sk.shape, np.int16); cls[ys, xs] = np.round(w_mm / BIN).astype(np.int16) + 1
    per = T.edge_lengths(sk, px, cls)
    return {str((c - 1) * BIN): round(float(v) * mm_pt / 1000, 3) for c, v in per.items() if c > 0}, sk


def scan_page(p, i, cfg, drs):
    text = p.get_text(); cnt = collections.Counter()
    for dr in drs:
        for name, col in cfg["colors"].items():
            if T.iscol(dr.get("color"), col, cfg["color_tol"]): cnt[name] += len(dr["items"])
    L = [l.strip() for l in text.split("\n") if l.strip()]
    gm = T.grid_scale(T.vertical_x(p, p.rect.width * cfg["content_xmax_frac"], drawings=drs), cfg["grid_gaps_mm"]) if cnt else None
    title = max((l for l in L if any(w in l.upper() for w in KEYWORDS)), key=len, default="")
    return {"sheet": i, "role": cfg["roles"].get(str(i), T.guess_role(text)), "items": dict(cnt),
            "label_scales": dict(T.label_scales(text)), "geometry_scale": T.snap_scale(gm[0]) if gm else None,
            "geometry_lines": (gm[1] + 1) if gm else 0, "supports": T.parse_supports(text), "legend": T.legend_widths(text),
            "is_index": "DRAWING INDEX" in text.upper(), "title": title[:120]}


def measure_page(p, row, cfg, drs, px, out):
    role = row["role"]; N, st, fl = T.resolve_scale(row, cfg)
    rec = {"sheet": row["sheet"], "role": role, "scale": N, "scale_status": st, "tray_hist_m": {}, "conduit_h_m": 0.0, "conduit_v_m": 0.0,
           "flags": [fl] if fl else []}
    if len(row["label_scales"]) > 1: rec["flags"].append(f"beberapa skala di satu sheet {sorted(row['label_scales'])}; dipakai: 1:{N}")
    if N is None: rec["flags"].append("dilewati: skala tidak diketahui"); return rec
    k = N * T.PT_MM; xmax = p.rect.width * cfg["content_xmax_frac"]
    if row["items"].get("tray") and role in cfg["tray_from"]:
        m = T.make_mask(p, cfg["colors"]["tray"], cfg["color_tol"], px, xmax, drawings=drs)
        h, sk = hist_measure(m, cfg["close_mm"]["tray"] / k, px, k)
        rec["tray_hist_m"] = h
        T.write_overlay(m, sk, os.path.join(out, "overlay", f"sheet{row['sheet']:03d}_tray.png"))
    if row["items"].get("conduit"):
        if role in cfg["conduit_horizontal_from"]:
            m = T.make_mask(p, cfg["colors"]["conduit"], cfg["color_tol"], px, xmax, drawings=drs)
            tot, _, _ = T.measure(m, cfg["close_mm"]["conduit"] / k, px); rec["conduit_h_m"] = round(float(tot) * k / 1000, 2)
        if role in cfg["conduit_vertical_from"]:
            m = T.make_mask(p, cfg["colors"]["conduit"], cfg["color_tol"], px, xmax, "v", drs)
            tot, _, _ = T.measure(m, cfg["close_mm"]["conduit"] / k, px); rec["conduit_v_m"] = round(float(tot) * k / 1000, 2)
    return rec


def run(pdf, cfg, out, budget=240, px=4):
    import pymupdf as fitz
    doc = fitz.open(pdf); n = len(doc); os.makedirs(os.path.join(out, "overlay"), exist_ok=True)
    path = os.path.join(out, "pages.jsonl"); done = set()
    if os.path.exists(path):
        for ln in open(path, encoding="utf-8"):
            try: done.add(json.loads(ln)["row"]["sheet"])
            except Exception: pass
    t0 = time.time()
    with open(path, "a", encoding="utf-8") as fh:
        for i in range(1, n + 1):
            if i in done: continue
            if time.time() - t0 > budget: break
            p = doc[i - 1]; drs = p.get_cdrawings(); row = scan_page(p, i, cfg, drs)
            rec = measure_page(p, row, cfg, drs, px, out) if (row["items"] and row["role"] in ("plan", "section")) else None
            fh.write(json.dumps({"row": row, "rec": rec}, ensure_ascii=False) + "\n"); fh.flush(); done.add(i)
    print(f"Progres: {len(done)}/{n} halaman ({len(done) * 100 // n}%).")
    if len(done) < n: print("BELUM selesai: jalankan perintah yang sama lagi untuk melanjutkan (hasil tersimpan di pages.jsonl)."); return None
    res = aggregate(out, cfg, os.path.basename(pdf)); print_sets(res); return res


def set_name(rows):
    for r in rows:
        if r["title"]: return re.split(r"ACC\w+\s+", r["title"], flags=re.I)[-1].strip(" -")
    return "(tanpa nama)"


def aggregate(out, cfg, fname=""):
    recs = [json.loads(l) for l in open(os.path.join(out, "pages.jsonl"), encoding="utf-8")]; recs.sort(key=lambda r: r["row"]["sheet"])
    n = len(recs); idx = [r["row"]["sheet"] for r in recs if r["row"]["is_index"]]
    starts = sorted({max(1, i - 1) for i in idx}) or [1]
    if starts[0] != 1: starts = [1] + starts
    letters = {k.upper(): (v, 0) for k, v in cfg["support_letter_width_mm"].items()}
    sets = []
    for si, a in enumerate(starts):
        b = (starts[si + 1] - 1) if si + 1 < len(starts) else n
        part = [r for r in recs if a <= r["row"]["sheet"] <= b]
        sup = {}; leg = dict(letters)
        for r in part: sup.update(r["row"]["supports"]); leg.update(r["row"]["legend"])
        sup_h = {k: v for k, v in sup.items() if k.startswith("SUH")}
        cand = sorted({leg[v["letter"]][0] for v in sup_h.values() if v["letter"] in leg}) or list(cfg["tray_widths_mm"])
        tray = collections.defaultdict(float)
        for r in part:
            for wb, v in ((r["rec"] or {}).get("tray_hist_m") or {}).items():
                tray[str(min(cand, key=lambda c: abs(c - int(wb))))] += v
        ch = sum((r["rec"] or {}).get("conduit_h_m", 0) for r in part); cv = sum((r["rec"] or {}).get("conduit_v_m", 0) for r in part)
        tray = {w: round(v, 1) for w, v in tray.items()}
        flags = []; lab_only = [r["row"]["sheet"] for r in part if r["rec"] and r["rec"]["scale_status"] == "label-saja"]
        for r in part:
            for f in (r["rec"] or {}).get("flags", []):
                if "TIDAK cocok" in f or "WAJIB" in f or "beberapa skala" in f or "tidak diketahui" in f: flags.append(f"hal {r['row']['sheet']}: {f}")
        unmeasured = [r["row"]["sheet"] for r in part if r["row"]["items"] and r["row"]["role"] == "other"]
        if unmeasured: flags.append(f"{len(unmeasured)} halaman berwarna TIDAK terukur (jenis halaman tidak dikenali): hal {unmeasured[:12]}")
        if lab_only: flags.append(f"{len(lab_only)} halaman berskala dari label saja (belum terverifikasi geometri)")
        if (ch or cv or tray) and not sup_h: flags.append("tidak ada qty support terbaca; cek silang spacing dilewati; lebar tray pakai config")
        name = set_name([r["row"] for r in part])
        sets.append({"set": si + 1, "pages": [a, b], "name": name, "tray_m": tray, "conduit_h_m": round(ch, 1), "conduit_v_upper_m": round(cv, 1),
                     "supports": {k: v["qty"] for k, v in sup_h.items()}, "support_total": sum(v["qty"] for v in sup_h.values()),
                     "checks": T.crosscheck(tray, sup_h, leg, cfg) if tray and sup_h else [], "flags": flags, "key": re.sub(r"[^A-Z0-9]", "", SKIP_WORDS.sub("", name.upper())), "include": True})
    seen = collections.defaultdict(list)
    for s in sets: seen[s["key"]].append(s)
    for key, grp in seen.items():
        if len(grp) > 1 and key:
            for g in grp[:-1]: g["include"] = False; g["flags"].append(f"DUPLIKAT: judul sama muncul {len(grp)}x di PDF (set {', '.join(str(x['set']) for x in grp)}); default hanya set terakhir dihitung, KONFIRMASI revisi yang benar")
            grp[-1]["flags"].append(f"versi terakhir dari {len(grp)} set berjudul sama (default dihitung)")
    res = {"file": fname, "source": "pdf-batch", "pages": n, "sets": sets}
    with open(os.path.join(out, "plan.json"), "w", encoding="utf-8") as fh: json.dump(res, fh, ensure_ascii=False, indent=1)
    return res


def print_sets(res):
    print("\n## DRY-RUN takeoff batch (belum ada Excel)\n")
    print("| Set | Halaman | Nama | Dihitung | Tray m (lebar mm) | Conduit horiz m | Conduit vert m (batas atas) | Support |\n|---|---|---|---|---|---|---|---|")
    tt = collections.defaultdict(float); ch = cv = 0
    for s in res["sets"]:
        tr = ", ".join(f"{w}W={v}" for w, v in sorted(s["tray_m"].items(), key=lambda x: int(x[0]))) or "-"
        print(f"| {s['set']} | {s['pages'][0]}-{s['pages'][1]} | {s['name'][:48]} | {'ya' if s['include'] else 'TIDAK (duplikat)'} | {tr} | {s['conduit_h_m'] or '-'} | {s['conduit_v_upper_m'] or '-'} | {s['support_total'] or '-'} |")
        if s["include"]:
            for w, v in s["tray_m"].items(): tt[w] += v
            ch += s["conduit_h_m"]; cv += s["conduit_v_upper_m"]
    print(f"\nTotal (set 'ya'): tray {dict(sorted(((w, round(v, 1)) for w, v in tt.items()), key=lambda x: int(x[0])))} | conduit horizontal {round(ch, 1)} m | conduit vertikal batas atas {round(cv, 1)} m")
    print("\nPeringatan:")
    for s in res["sets"]:
        for f in s["flags"][:6]: print(f" - set {s['set']} ({s['name'][:30]}): {f}")
        for c in s["checks"]: print(f"   cek support set {s['set']}: {c}")


def area_label(name):
    """Judul set -> label area singkat: 'A2O BIOREACTOR TANKS (NORTH)' -> 'A2O North'."""
    n = re.sub(r"\s*\((NORTH|SOUTH)\)", r" \1", name.strip(), flags=re.I).title()
    n = re.sub(r"\bBioreactor Tanks\b", "", n); n = re.sub(r"\bTanks\b", "", n) if n.startswith("Mbr") else n
    n = re.sub(r"\bMbr\b", "MBR", n); n = re.sub(r"\bWas\b", "WAS", n); n = n.replace("_", " / ")
    return re.sub(r"\s+", " ", n).strip()


def reliability(s):
    f = " ".join(s["flags"]); ch = " ".join(s["checks"])
    if "PERINGATAN" in ch: return "TIDAK andal: cek support gagal"
    if (s["support_total"] or s["tray_m"]) and not s["conduit_h_m"] and not s["conduit_v_upper_m"]: return "Conduit kosong: cek manual"
    if "TIDAK terukur" in f: return "Sebagian halaman tidak terukur"
    if not s["checks"] and (s["tray_m"] or s["conduit_h_m"]): return "Belum terverifikasi (tanpa cek support)"
    return "Cek support OK" if s["checks"] else "Tanpa data"


def write_excel(out):
    """Rekap per area (baris) x ukuran (kolom): tray per lebar, conduit total. Rumus SUMIFS ke sheet 'Detail Set' (kolom Hitung Y/N)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter as CL
    res = json.load(open(os.path.join(out, "plan.json"), encoding="utf-8")); sets = res["sets"]
    found = {int(w) for s in sets for w in s["tray_m"]}
    first = [w for w in (700, 600, 500) if True]; rest = sorted((w for w in found if w not in first), reverse=True)
    widths = first + rest
    for s in sets: s["area"] = area_label(s["name"]); s["rel"] = reliability(s)
    areas = []
    for a in ["A2O North", "A2O South", "Administration Building"] + [s["area"] for s in sets]:
        if a not in areas and any(s["area"] == a for s in sets): areas.append(a)
    F = Font(name="Arial"); FB = Font(name="Arial", bold=True); HD = PatternFill("solid", fgColor="305496"); GR = PatternFill("solid", fgColor="D9D9D9")
    RED = PatternFill("solid", fgColor="F8CBAD"); thin = Side(style="thin", color="999999"); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
    wb = Workbook(); wr = wb.active; wr.title = "Rekap per Area"; wd = wb.create_sheet("Detail Set")
    # ---- Detail Set
    dh = ["Set", "Area", "Hal awal", "Hal akhir", "Hitung (Y/N)"] + [f"Tray {w}" for w in widths] + ["Conduit horizontal (m)", "Conduit vertikal batas atas (m)", "Support (EA)", "Keandalan", "Catatan"]
    wd.append(dh)
    for c in wd[1]: c.font = Font(name="Arial", bold=True, color="FFFFFF"); c.fill = HD; c.alignment = Alignment(wrap_text=True, vertical="center")
    for s in sets:
        wd.append([s["set"], s["area"], s["pages"][0], s["pages"][1], "Y" if s["include"] else "N"] + [s["tray_m"].get(str(w), 0) for w in widths] +
                  [s["conduit_h_m"], s["conduit_v_upper_m"], s["support_total"], s["rel"], "; ".join(s["flags"])[:700]])
    nd = wd.max_row; col = {w: CL(6 + i) for i, w in enumerate(widths)}; ch_c = CL(6 + len(widths)); cv_c = CL(7 + len(widths))
    for i in range(1, wd.max_column + 1): wd.column_dimensions[CL(i)].width = 80 if i == wd.max_column else (30 if i in (2, wd.max_column - 1) else 12)
    wd.freeze_panes = "F2"
    for r in wd.iter_rows(min_row=2):
        for c in r: c.font = F
    # ---- Rekap
    n_t = len(widths); cs = ["16", "22", "28"]
    wr["A1"] = "Rekap Quantity Cable Tray dan Conduit per Area (meter)"; wr["A1"].font = Font(name="Arial", bold=True, size=13)
    r0 = 3; wr.cell(r0, 1, "")
    wr.cell(r0, 2, "Cable Tray (lebar mm)"); wr.merge_cells(start_row=r0, start_column=2, end_row=r0, end_column=1 + n_t)
    wr.cell(r0, 2 + n_t, "Conduit (ukuran mm)"); wr.merge_cells(start_row=r0, start_column=2 + n_t, end_row=r0, end_column=4 + n_t + 1)
    for j, w in enumerate(widths): wr.cell(r0 + 1, 2 + j, w)
    for j, x in enumerate(cs): wr.cell(r0 + 1, 2 + n_t + j, int(x))
    wr.cell(r0 + 1, 2 + n_t + 3, "Total horizontal"); wr.cell(r0 + 1, 2 + n_t + 4, "Vertikal (batas atas)")
    last_c = 2 + n_t + 4 + 3
    for j, h in enumerate(["Set dihitung (hal.)", "Keandalan", "Catatan"]): wr.cell(r0 + 1, 3 + n_t + 4 + j, h)
    for rr in (r0, r0 + 1):
        for cc in range(1, last_c + 1):
            c = wr.cell(rr, cc); c.font = Font(name="Arial", bold=True, color="FFFFFF"); c.fill = HD; c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); c.border = BOX
    wr.cell(r0 + 1, 1, "Area").alignment = Alignment(horizontal="left")
    for i, a in enumerate(areas):
        r = r0 + 2 + i; wr.cell(r, 1, a)
        for j, w in enumerate(widths):
            wr.cell(r, 2 + j, f"=SUMIFS('Detail Set'!${col[w]}$2:${col[w]}${nd},'Detail Set'!$B$2:$B${nd},$A{r},'Detail Set'!$E$2:$E${nd},\"Y\")").number_format = '0.0;-0.0;"-"'
        for j in range(3):
            c = wr.cell(r, 2 + n_t + j, "belum diukur"); c.font = Font(name="Arial", italic=True, color="7F7F7F", size=9); c.fill = GR; c.alignment = Alignment(horizontal="center")
        for j, cc in enumerate((ch_c, cv_c)):
            wr.cell(r, 2 + n_t + 3 + j, f"=SUMIFS('Detail Set'!${cc}$2:${cc}${nd},'Detail Set'!$B$2:$B${nd},$A{r},'Detail Set'!$E$2:$E${nd},\"Y\")").number_format = '0.0;-0.0;"-"'
        inc = [s for s in sets if s["area"] == a and s["include"]]; s = inc[-1] if inc else None
        wr.cell(r, 3 + n_t + 4, f"set {s['set']}: {s['pages'][0]}-{s['pages'][1]}" if s else "tidak ada yang dihitung")
        wr.cell(r, 4 + n_t + 4, s["rel"] if s else ""); wr.cell(r, 5 + n_t + 4, "; ".join(s["flags"])[:300] if s else "")
        bad = bool(s and ("TIDAK andal" in s["rel"] or "kosong" in s["rel"]))
        for cc in range(1, last_c + 1):
            c = wr.cell(r, cc); c.border = BOX
            if not (2 + n_t <= cc < 2 + n_t + 3): c.font = F
            if bad and not (2 + n_t <= cc < 2 + n_t + 3): c.fill = RED
    rt = r0 + 2 + len(areas); wr.cell(rt, 1, "TOTAL")
    for cc in list(range(2, 2 + n_t)) + [2 + n_t + 3, 2 + n_t + 4]:
        wr.cell(rt, cc, f"=SUM({CL(cc)}{r0 + 2}:{CL(cc)}{rt - 1})").number_format = '0.0;-0.0;"-"'
    for cc in range(1, last_c + 1): wr.cell(rt, cc).font = FB; wr.cell(rt, cc).border = BOX
    wr.column_dimensions["A"].width = 30
    for cc in range(2, last_c + 1): wr.column_dimensions[CL(cc)].width = 12
    wr.column_dimensions[CL(3 + n_t + 4)].width = 20; wr.column_dimensions[CL(4 + n_t + 4)].width = 34; wr.column_dimensions[CL(5 + n_t + 4)].width = 90
    wr.freeze_panes = wr.cell(r0 + 2, 2)
    notes = ["Catatan:", "1. Hanya cable tray dan conduit. Ukur vektor PDF, akurasi sekitar +/-5%, belum termasuk waste.",
             "2. Tray per lebar dari lebar garis (dipetakan ke lebar support di set); lebar selain 700/600/500 tetap ditampilkan karena terukur.",
             "3. Conduit per ukuran (16/22/28) BELUM bisa: garis hijau tidak membedakan diameter. Yang tersedia: total horizontal (plan) dan vertikal (batas atas, belum dedup antar-section).",
             "4. Revisi/duplikat diatur di sheet 'Detail Set' kolom Hitung (Y/N); rekap terhitung ulang otomatis.",
             "5. Baris berwarna merah muda = hasil tidak andal (cek support gagal atau conduit kosong); jangan dipakai sebelum dicek manual."]
    for i, t in enumerate(notes): wr.cell(rt + 2 + i, 1, t).font = FB if i == 0 else F
    path = os.path.join(out, "Rekap_Tray_Conduit_per_Area.xlsx"); wb.save(path); print("Excel:", path)

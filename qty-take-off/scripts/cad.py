#!/usr/bin/env python3
"""Jalur CAD untuk qty-take-off: DWG/DXF -> panjang tray/conduit EKSAK dari entitas (tanpa skala, tanpa raster).

  python cad.py convert FILE.dwg --out DIR     # DWG -> DXF (dwg2dxf / ODAFileConverter), verifikasi terbaca
  python cad.py scan    FILE.dxf               # layer, satuan, layout, blok support, kandidat layer tray/conduit
Dipakai oleh takeoff.py (plan/execute) bila input .dwg/.dxf. Hasil lebih akurat dari PDF; PDF = cadangan."""
import collections, difflib, json, math, os, re, shutil, subprocess, sys, tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import cadgeom as G

SUPPORT_RE = re.compile(r"^SUH-\d+([A-K])[A-Z]$", re.I)
TRAY_RE = re.compile(r"tray|ladder|cabletray|\bct[-_ ]|trunk", re.I)
CONDUIT_RE = re.compile(r"conduit|cndt|\bcdt\b|g-type|gtype", re.I)
ACI = {"tray": 6, "conduit": 3}          # magenta, hijau (cadangan bila nama layer tidak informatif)
VALID_CONDUIT = (16, 22, 28, 36, 42, 54, 70, 82, 92, 104)
TYPE_WORDS = {"POWER": "POWER", "PWR": "POWER", "SIGNAL": "SIGNAL", "SIG": "SIGNAL",
              "CONTROL": "CONTROL", "CTRL": "CONTROL", "CTL": "CONTROL"}
MAX_RUNS = 300


class ConvertError(Exception): pass
class CadGeometryError(ValueError): pass


# ------------------------------------------------------------ konversi DWG -> DXF
def dwg_version(path):
    with open(path, "rb") as f: tag = f.read(6).decode("ascii", "ignore")
    return {"AC1032": "AutoCAD 2018+", "AC1027": "AutoCAD 2013", "AC1024": "AutoCAD 2010", "AC1021": "AutoCAD 2007",
            "AC1018": "AutoCAD 2004", "AC1015": "AutoCAD 2000"}.get(tag, f"tidak dikenal ({tag})")


def find_converter():
    for name, env in (("dwg2dxf", "DWG2DXF_BIN"), ("ODAFileConverter", "ODA_BIN")):
        p = os.environ.get(env) or shutil.which(name)
        if p and os.path.exists(p): return name, p
    import convert_local
    p = convert_local.find_oda()
    if p: return "ODAFileConverter", p
    return None, None


def convert_dwg(path, out_dir, timeout=280):
    """Return path DXF. Gagal -> ConvertError dengan instruksi yang bisa langsung disampaikan ke user."""
    os.makedirs(out_dir, exist_ok=True)
    ver = dwg_version(path); size = os.path.getsize(path) / 1e6
    kind, bin_ = find_converter()
    hint = (f"DWG {ver}, {size:.0f} MB tidak dapat dikonversi otomatis di sini. Minta user mengekspor dari AutoCAD: "
            "SAVEAS -> DXF (versi 2013/2018 ASCII), atau PDF vektor (PLOT, bukan scan), lalu unggah ulang. "
            "Alternatif: konversi di laptop dengan scripts/convert_local.py (ODA File Converter), lalu unggah DXF-nya.")
    if not kind: raise ConvertError("Tidak ada converter DWG (dwg2dxf/ODAFileConverter); coba `bash scripts/setup_converter.sh` sekali. " + hint)
    out = os.path.join(out_dir, os.path.splitext(os.path.basename(path))[0] + ".dxf")
    try:
        if kind == "dwg2dxf":
            r = subprocess.run([bin_, "-y", "-o", out, path], capture_output=True, text=True, timeout=timeout,
                               env={**os.environ, "LD_LIBRARY_PATH": os.path.join(os.path.dirname(os.path.dirname(bin_)), "lib") + ":" + os.environ.get("LD_LIBRARY_PATH", "")})
        else:                              # ODAFileConverter in_dir out_dir ACAD2018 DXF 0 1 filter
            with tempfile.TemporaryDirectory() as src:
                shutil.copy(path, src)
                r = subprocess.run([bin_, src, out_dir, "ACAD2018", "DXF", "0", "1", os.path.basename(path)],
                                   capture_output=True, text=True, timeout=timeout)
    except subprocess.TimeoutExpired:
        raise ConvertError(f"Konversi melebihi {timeout} detik. " + hint)
    if not os.path.exists(out) or os.path.getsize(out) < 1000:
        tail = (r.stderr or r.stdout or "").strip().splitlines()[-2:]
        raise ConvertError(f"{kind} gagal ({' | '.join(tail)}). " + hint)
    try: load(out)
    except Exception as e: raise ConvertError(f"DXF hasil konversi tidak terbaca ({type(e).__name__}). " + hint)
    return out


# ------------------------------------------------------------ baca DXF
def load(path):
    import ezdxf
    return ezdxf.readfile(path)


def unit_factor_to_m(doc, assume="mm"):
    """Faktor satuan gambar -> meter. INSUNITS=0 (unitless) -> asumsi config + peringatan."""
    from ezdxf import units
    u = doc.units
    if u == 0:
        return units.conversion_factor(getattr(units, assume.upper(), units.MM), units.M), f"INSUNITS tidak diisi; diasumsikan {assume}"
    return units.conversion_factor(u, units.M), None


def _flat_len(e):
    """Panjang entitas melengkung/polyline lewat flattening path (tetap 3D bila ada z)."""
    from ezdxf import path as ep
    pts = list(ep.make_path(e).flattening(0.01))
    return [(pts[i], pts[i + 1]) for i in range(len(pts) - 1)]


def segments(e):
    """Entitas -> daftar segmen lurus ((x,y,z),(x,y,z)). Tipe tak dikenal -> []."""
    t = e.dxftype()
    if t == "LINE": return [(tuple(e.dxf.start), tuple(e.dxf.end))]
    if t in ("LWPOLYLINE", "POLYLINE", "ARC", "SPLINE", "ELLIPSE"):
        try:
            result = [(tuple(a), tuple(b)) for a, b in _flat_len(e)]
        except Exception as exc:
            raise CadGeometryError(f"Gagal membaca geometri {t} handle={e.dxf.get('handle', '?')}: {exc}") from exc
        if not result:
            raise CadGeometryError(f"Geometri {t} handle={e.dxf.get('handle', '?')} tidak menghasilkan segmen.")
        return result
    return []


def walk(layout, depth=0, parent_layer=None, max_depth=4):
    """Hasil: (entitas, layer_efektif). INSERT dibongkar rekursif; layer '0' di dalam blok mewarisi layer INSERT."""
    for e in layout:
        t = e.dxftype()
        lay = e.dxf.layer if e.dxf.layer != "0" or parent_layer is None else parent_layer
        if t == "INSERT" and depth >= max_depth:
            raise CadGeometryError(f"Blok bertingkat melebihi batas {max_depth}: {e.dxf.name} handle={e.dxf.get('handle', '?')}")
        if t == "INSERT":
            try:
                for v in e.virtual_entities():
                    yield from walk([v], depth + 1, e.dxf.layer if e.dxf.layer != "0" else parent_layer)
            except Exception as exc:
                raise CadGeometryError(f"Gagal membongkar blok {e.dxf.name} handle={e.dxf.get('handle', '?')}: {exc}") from exc
        elif t != "INSERT": yield e, lay


def dist3(a, b): return math.dist(a, b)


def hv(a, b):
    h = math.hypot(b[0] - a[0], b[1] - a[1]); v = abs(b[2] - a[2]); return h, v


def conduit_cat(layer_name):
    """Return the (size, type) only when both are encoded in a layer name."""
    upper = (layer_name or "").upper()
    size = next((int(value) for value in re.findall(r"(?<!\d)(\d{2,3})(?!\d)", upper)
                 if int(value) in VALID_CONDUIT), None)
    kind = next((value for word, value in TYPE_WORDS.items()
                 if re.search(r"(?<![A-Z])%s(?![A-Z])" % word, upper)), None)
    return (size, kind) if size and kind else None


def parse_label_text(text):
    """Parse size/type/elevation from one CAD text entity, when available."""
    match = re.search(r"(?<!\d)(\d{2,3})\s*mm", text or "", re.I)
    if not match:
        return None
    remainder = text[match.end():]
    kind = None
    for token in re.findall(r"[A-Za-z]{3,}", remainder):
        upper = token.upper()
        if upper in TYPE_WORDS:
            kind = TYPE_WORDS[upper]
            break
        if len(upper) >= 4:
            best = max(("POWER", "SIGNAL", "CONTROL"),
                       key=lambda value: difflib.SequenceMatcher(None, upper, value).ratio())
            if difflib.SequenceMatcher(None, upper, best).ratio() >= 0.72:
                kind = best
                break
    if not kind:
        return None
    elevation = re.search(r"EL\s*[+-]?\s*(\d{3,6})", remainder, re.I)
    return int(match.group(1)), kind, int(elevation.group(1)) if elevation else None


def _inside(window, x, y):
    return window[0] <= x <= window[2] and window[1] <= y <= window[3]


def _validate_windows(windows, name):
    for window in windows:
        if (len(window) != 4 or window[0] > window[2] or window[1] > window[3]):
            raise ValueError(f"{name} harus berisi [x_min, y_min, x_max, y_max] yang terurut.")


def _clip_segment(start, end, include, exclude):
    """Split a segment at window edges and retain only points inside configured windows."""
    if not include and not exclude:
        return [(start, end)]
    delta_x, delta_y = end[0] - start[0], end[1] - start[1]
    parameters = {0.0, 1.0}
    for window in (*include, *exclude):
        if delta_x:
            parameters.update((value - start[0]) / delta_x for value in (window[0], window[2])
                              if 0 < (value - start[0]) / delta_x < 1)
        if delta_y:
            parameters.update((value - start[1]) / delta_y for value in (window[1], window[3])
                              if 0 < (value - start[1]) / delta_y < 1)
    ordered = sorted(parameters)
    pieces = []
    for first, last in zip(ordered, ordered[1:]):
        middle = (first + last) / 2
        x, y = start[0] + middle * delta_x, start[1] + middle * delta_y
        inside_include = not include or any(_inside(window, x, y) for window in include)
        if inside_include and not any(_inside(window, x, y) for window in exclude):
            pieces.append((
                tuple(start[i] + first * (end[i] - start[i]) for i in range(3)),
                tuple(start[i] + last * (end[i] - start[i]) for i in range(3)),
            ))
    return pieces


# ------------------------------------------------------------ scan
def scan(path, cfg=None):
    doc = load(path); msp = doc.modelspace(); k, uw = unit_factor_to_m(doc, (cfg or {}).get("assume_units", "mm"))
    lay = collections.defaultdict(lambda: collections.Counter()); length = collections.defaultdict(float)
    box = {}; sup = collections.Counter(); has_z = False
    for e, l in walk(msp):
        lay[l][e.dxftype()] += 1
        for a, b in segments(e):
            length[l] += dist3(a, b) * k
            if abs(a[2]) > 1e-6 or abs(b[2]) > 1e-6: has_z = True
            bounds = box.setdefault(l, [a[0], a[1], a[0], a[1]])
            for point in (a, b):
                bounds[0] = min(bounds[0], point[0]); bounds[1] = min(bounds[1], point[1])
                bounds[2] = max(bounds[2], point[0]); bounds[3] = max(bounds[3], point[1])
    for e in msp:
        if e.dxftype() == "INSERT" and SUPPORT_RE.match(e.dxf.name): sup[e.dxf.name.upper()] += 1
    layers = []
    for name, cnt in sorted(lay.items()):
        ent = doc.layers.get(name) if name in doc.layers else None
        role = "tray" if TRAY_RE.search(name) else "conduit" if CONDUIT_RE.search(name) else None
        color = ent.color if ent is not None else None
        if role is None and color in (ACI["tray"], ACI["conduit"]): role = {6: "tray", 3: "conduit"}[color] + "?(warna)"
        layers.append({"layer": name, "entities": dict(cnt), "length_m": round(length[name], 1), "aci": color, "role_guess": role,
                       "bbox_du": [round(value, 1) for value in box[name]] if name in box else None})
    layouts = [{"name": ly.name, "viewports": len([v for v in ly if v.dxftype() == "VIEWPORT"])} for ly in doc.layouts if ly.name != "Model"]
    return {"file": os.path.basename(path), "dxfversion": doc.dxfversion, "units": int(doc.units), "units_warning": uw,
            "has_3d": has_z, "layers": layers, "layouts": layouts, "support_blocks": dict(sup), "xrefs": [b.name for b in doc.blocks if b.block.dxf.get("flags", 0) & 4]}


def suggest_config(scan_result):
    """Create a user-reviewed layer config from a CAD scan."""
    tray = [layer["layer"] for layer in scan_result["layers"]
            if (layer["role_guess"] or "").startswith("tray")]
    conduit = [layer["layer"] for layer in scan_result["layers"]
               if (layer["role_guess"] or "").startswith("conduit")]
    pattern = lambda names: "^(?:" + "|".join(re.escape(name) for name in names) + ")$" if names else "(?!x)x"
    widths = {}
    for name in tray:
        match = re.search(r"(?<!\d)(\d{3,4})(?!\d)", name)
        widths[name] = int(match.group(1)) if match else None
    config = {
        "layers": {"tray": pattern(tray), "conduit": pattern(conduit)},
        "layer_widths": widths,
        "include_windows": [],
        "exclude_windows": [],
        "_todo": [
            "Periksa daftar layer tray/conduit.",
            "Isi layer_widths yang null (mm).",
            "Jika plan dan section berdampingan di model space, isi include_windows "
            "[[x_min,y_min,x_max,y_max], ...] agar hanya view plan yang dihitung. "
            "Koordinat gambar ada di bbox_du hasil scan.",
        ],
    }
    if conduit and not any(conduit_cat(name) for name in conduit):
        config["_todo"].append("Nama layer conduit tidak memuat ukuran+jenis; rincian kategori hanya dari teks CAD.")
    return config


# ------------------------------------------------------------ ukur
def layer_width_mm(name, cfg):
    if name in cfg.get("layer_widths", {}): return cfg["layer_widths"][name]
    m = re.search(r"(?<!\d)(\d{3,4})(?!\d)", name)
    if m and int(m.group(1)) in cfg["tray_widths_mm"]: return int(m.group(1))
    return None


def measure(path, cfg):
    """-> dict hasil seragam dengan jalur PDF (satu 'sheet' bernama 'Model')."""
    doc = load(path); msp = doc.modelspace(); k, uw = unit_factor_to_m(doc, cfg.get("assume_units", "mm"))
    tray_re = re.compile(cfg.get("layers", {}).get("tray", TRAY_RE.pattern), re.I)
    cond_re = re.compile(cfg.get("layers", {}).get("conduit", CONDUIT_RE.pattern), re.I)
    include = cfg.get("include_windows") or []
    exclude = cfg.get("exclude_windows") or []
    _validate_windows(include, "include_windows"); _validate_windows(exclude, "exclude_windows")
    tolerance = cfg.get("merge_tol_mm", 2) / 1000.0
    link_tolerance = cfg.get("run_link_mm", 50) / 1000.0
    unsupported = collections.Counter()
    raw = collections.defaultdict(list)
    cut = collections.defaultdict(float)
    labels = []
    has_z = False
    flags = []
    if uw: flags.append(uw)
    for layer in doc.layers:
        name = layer.dxf.name
        if layer.dxf.color == ACI["tray"] and not tray_re.search(name):
            flags.append(f"layer {name} berwarna magenta tetapi belum dipetakan sebagai tray; cek hasil scan")
        elif layer.dxf.color == ACI["conduit"] and not cond_re.search(name):
            flags.append(f"layer {name} berwarna hijau tetapi belum dipetakan sebagai conduit; cek hasil scan")
    for e, l in walk(msp):
        entity_type = e.dxftype()
        if entity_type in ("TEXT", "MTEXT"):
            text = e.plain_text() if entity_type == "MTEXT" else e.dxf.text
            position = e.dxf.insert
            if (not include or any(_inside(window, position[0], position[1]) for window in include)) and not any(
                    _inside(window, position[0], position[1]) for window in exclude):
                label = parse_label_text(text or "")
                if label:
                    labels.append(label)
        role = "tray" if tray_re.search(l) else "conduit" if cond_re.search(l) else None
        if role is None: continue
        segs = segments(e)
        if not segs:
            unsupported[(role, entity_type)] += 1
            continue
        category = layer_width_mm(l, cfg) if role == "tray" else conduit_cat(l)
        for a, b in segs:
            kept = _clip_segment(a, b, include, exclude)
            if any(abs(first[2] - last[2]) > 1e-6 for first, last in kept):
                has_z = True
            cut[role] += max(0.0, dist3(a, b) - sum(dist3(first, last) for first, last in kept)) * k
            raw[(role, category)].extend(
                (tuple(coord * k for coord in first), tuple(coord * k for coord in last))
                for first, last in kept
            )
    merged = {}
    duplicates = collections.defaultdict(float)
    for key, segments_ in raw.items():
        before = sum(dist3(a, b) for a, b in segments_)
        unique = G.merge_collinear(segments_, tolerance)
        duplicates[key[0]] += max(0.0, before - sum(dist3(a, b) for a, b in unique))
        merged[key] = unique

    width_max = max(cfg.get("tray_widths_mm") or [0]) / 1000.0
    rung_length = cfg.get("rung_factor", 1.5) * width_max
    tray_lengths = {}
    tray_orientations = {}
    runs = []
    for (role, width), segments_ in merged.items():
        if role != "tray":
            continue
        lengths = [dist3(a, b) for a, b in segments_]
        total = sum(lengths)
        short = sum(length for length in lengths if length <= rung_length)
        ladder = total > 0 and short / total > 0.15 and len(segments_) > 20
        usable = [segment for segment, length in zip(segments_, lengths)
                  if not ladder or length > rung_length]
        factor = 0.5 if ladder else 1.0
        key = str(width) if width else "tanpa-lebar"
        tray_lengths[key] = tray_lengths.get(key, 0.0) + sum(dist3(a, b) for a, b in usable) * factor
        orientation = G.orient_split(usable)
        totals = tray_orientations.setdefault(key, {"h": 0.0, "v": 0.0, "d": 0.0})
        for axis in "hvd":
            totals[axis] += orientation[axis] * factor
        if ladder:
            flags.append(f"tray lebar {width}: terdeteksi gambar ladder (rel+rung); panjang = rel/2; KONFIRMASI ke user")
        else:
            for run in G.cluster_runs(segments_, link_tolerance):
                runs.append({
                    "width": key,
                    "n": run["n"],
                    "length_m": round(run["length_m"], 2),
                    "h_m": round(run["h_m"], 2),
                    "v_m": round(run["v_m"], 2),
                    "d_m": round(run["d_m"], 2),
                    "bbox_du": [round(value / k, 1) for value in run["bbox"]],
                })
        if not width:
            flags.append("lebar tray tidak terbaca dari nama layer; isi layer_widths di config")
    tray_lengths = {width: round(length, 1) for width, length in tray_lengths.items()}
    runs.sort(key=lambda run: -run["length_m"])
    if len(runs) > MAX_RUNS:
        flags.append(f"{len(runs)} run tray; hanya {MAX_RUNS} terpanjang ditulis (sisanya fragmen kecil)")
        runs = runs[:MAX_RUNS]

    conduit_h = conduit_v = 0.0
    conduit_by_category = collections.defaultdict(lambda: {"h": 0.0, "v": 0.0})
    conduit_uncategorized = {"h": 0.0, "v": 0.0}
    for (role, category), segments_ in merged.items():
        if role != "conduit":
            continue
        for start, end in segments_:
            horizontal, vertical = hv(start, end)
            conduit_h += horizontal
            conduit_v += vertical
            target = conduit_by_category[category] if category else conduit_uncategorized
            target["h"] += horizontal
            target["v"] += vertical
    for role, length in duplicates.items():
        if length > 0.5:
            flags.append(f"{role}: {length:.1f} m garis ganda/tumpang tindih (geometri sama dalam toleransi) dibuang")
    for role, length in cut.items():
        if length > 0.5:
            flags.append(f"{role}: {length:.1f} m di luar include_windows / di dalam exclude_windows (tidak dihitung)")
    for (role, entity), count in unsupported.items():
        flags.append(f"{role}: {count} entitas {entity} pada layer terukur belum didukung; hasil mungkin tidak lengkap")
    if conduit_by_category and sum(conduit_uncategorized.values()) > 0.5:
        flags.append("conduit pada layer tanpa ukuran+jenis: "
                     f"horizontal {conduit_uncategorized['h']:.1f} m, vertikal {conduit_uncategorized['v']:.1f} m "
                     "(tidak masuk rincian per ukuran)")
    if not conduit_v:
        flags.append("tidak ada selisih Z pada geometri conduit; jalur vertikal perlu dihitung dari elevasi/callout")
    if not has_z and not include:
        flags.append("2D tanpa include_windows: view plan dan section di model space bisa TERCAMPUR (hitung ganda); "
                     "tentukan jendela per view (scan: bbox_du)")
    xrefs = [block.name for block in doc.blocks if block.block.dxf.get("flags", 0) & 4]
    if xrefs:
        flags.append(f"XREF tidak dihitung: {', '.join(xrefs)}")
    layouts = [layout.name for layout in doc.layouts if layout.name != "Model"]
    if layouts:
        flags.append(f"paper space tidak dihitung: {', '.join(layouts)}")
    labels_ok = collections.Counter(f"{size}|{kind}" for size, kind, _ in labels if size in VALID_CONDUIT)
    labels_excluded = collections.Counter(size for size, _, _ in labels if size not in VALID_CONDUIT)
    if labels_excluded:
        flags.append(f"label teks ukuran di luar daftar conduit {sorted(labels_excluded)}: "
                     "kabel pada tray (>104 mm) atau salah ketik; tidak dihitung")
    sup = collections.Counter(entity.dxf.name.upper() for entity in msp
                               if entity.dxftype() == "INSERT" and SUPPORT_RE.match(entity.dxf.name))
    letters = cfg.get("support_letter_width_mm", {})
    supports = {name: {"qty": quantity, "letter": SUPPORT_RE.match(name).group(1).upper()}
                for name, quantity in sup.items()}
    record = {
        "sheet": "Model", "role": "model", "scale": None,
        "scale_status": "CAD (satuan gambar, tanpa skala)", "tray_m": tray_lengths,
        "conduit_h_m": round(conduit_h, 1), "conduit_v_m": round(conduit_v, 1),
        "flags": flags,
        "tray_orient_m": {width: {axis: round(value, 2) for axis, value in orientation.items()}
                          for width, orientation in tray_orientations.items()},
        "runs": runs,
        "conduit_by_cat": [
            {"size": size, "type": kind, "h_m": round(value["h"], 2), "v_m": round(value["v"], 2)}
            for (size, kind), value in sorted(
                conduit_by_category.items(),
                key=lambda item: (item[0][0] is None, item[0][0] or 0, item[0][1] or ""),
            ) if size and kind
        ],
        "conduit_uncat_m": {axis: round(value, 2) for axis, value in conduit_uncategorized.items()},
        "labels_text": dict(labels_ok),
        "labels_excluded": {str(size): count for size, count in labels_excluded.items()},
    }
    return record, supports, {key.upper(): (value, 0) for key, value in letters.items()}


def print_extra(record):
    print("\nTray CAD per orientasi (m):")
    for width, orientation in sorted(record.get("tray_orient_m", {}).items()):
        print(f" - {width}: horizontal {orientation['h']}, vertikal {orientation['v']}, miring {orientation['d']}")
    if record.get("conduit_by_cat"):
        print("\nConduit CAD per ukuran x jenis dari nama layer:")
        print("| Ukuran | Jenis | Horizontal m | Vertikal m |\n|---|---|---:|---:|")
        for row in record["conduit_by_cat"]:
            print(f"| {row['size']} mm | {row['type'].title()} | {row['h_m']} | {row['v_m']} |")
    if record.get("labels_text"):
        print("\nInventaris label teks CAD (jumlah label, bukan meter): "
              + "; ".join(f"{key.replace('|', ' ')} x{count}"
                          for key, count in sorted(record["labels_text"].items())))
    if record.get("runs"):
        print("\nRun tray terpanjang (bbox dalam satuan gambar):")
        print("| Lebar | Panjang m | H | V | Miring | Segmen | bbox x0,y0,x1,y1 |\n|---|---:|---:|---:|---:|---:|---|")
        for run in record["runs"][:15]:
            print(f"| {run['width']} | {run['length_m']} | {run['h_m']} | {run['v_m']} | "
                  f"{run['d_m']} | {run['n']} | {run['bbox_du']} |")
        if len(record["runs"]) > 15:
            print(f"... {len(record['runs']) - 15} run lain ada di Excel (sheet Runs Tray).")


def add_cad_sheets(workbook, record):
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    def add_sheet(name, headers, rows):
        sheet = workbook.create_sheet(name)
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = Font(name="Arial", bold=True)
        for row in rows:
            sheet.append(row)
        for index in range(1, len(headers) + 1):
            sheet.column_dimensions[get_column_letter(index)].width = 22

    add_sheet(
        "Tray per Orientasi",
        ["Lebar (mm)", "Horizontal (m)", "Vertikal (m)", "Miring (m)"],
        [[width, orientation["h"], orientation["v"], orientation["d"]]
         for width, orientation in record.get("tray_orient_m", {}).items()],
    )
    add_sheet(
        "Conduit per Layer",
        ["Ukuran (mm)", "Jenis", "Horizontal (m)", "Vertikal (m)"],
        [[row["size"], row["type"], row["h_m"], row["v_m"]]
         for row in record.get("conduit_by_cat", [])]
        + [["tanpa kategori", "", record.get("conduit_uncat_m", {}).get("h", 0),
            record.get("conduit_uncat_m", {}).get("v", 0)]],
    )
    add_sheet(
        "Label Teks",
        ["Ukuran|Jenis", "Jumlah label"],
        [[key, count] for key, count in sorted(record.get("labels_text", {}).items())],
    )
    add_sheet(
        "Runs Tray",
        ["Lebar", "Panjang (m)", "H (m)", "V (m)", "Miring (m)", "Segmen",
         "x0", "y0", "x1", "y1 (satuan gambar)"],
        [[run["width"], run["length_m"], run["h_m"], run["v_m"], run["d_m"], run["n"]]
         + run["bbox_du"] for run in record.get("runs", [])],
    )


if __name__ == "__main__":
    if len(sys.argv) < 3: raise SystemExit(__doc__)
    cmd, f = sys.argv[1], sys.argv[2]
    if cmd == "convert":
        out = sys.argv[sys.argv.index("--out") + 1] if "--out" in sys.argv else "/tmp/tk"
        try: print(convert_dwg(f, out))
        except ConvertError as e: raise SystemExit(f"GAGAL: {e}")
    elif cmd == "scan": print(json.dumps(scan(f), ensure_ascii=False, indent=1))
    elif cmd == "suggest":
        output = sys.argv[sys.argv.index("--out") + 1] if "--out" in sys.argv else "cad_config.json"
        with open(output, "w", encoding="utf-8") as file:
            json.dump(suggest_config(scan(f)), file, ensure_ascii=False, indent=1)
        print("Template config ditulis:", output, "(periksa _todo sebelum dipakai)")

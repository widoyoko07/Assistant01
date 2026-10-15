"""Tes jalur CAD (DXF sintetis berpanjang diketahui). Jalankan: python3 -m unittest discover -s tests"""
import os, subprocess, sys, tempfile, unittest
sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "scripts"))
import ezdxf, cad, takeoff as T

def dxf(path, units=4, ladder=False, dup=True, z=True):
    d = ezdxf.new("R2018", setup=True); d.units = units
    for n in ("E-TRAY-700", "E-TRAY-200", "E-CONDUIT"): d.layers.add(n)
    m = d.modelspace()
    if ladder:
        m.add_lwpolyline([(0, 0), (10000, 0)], dxfattribs={"layer": "E-TRAY-700"}); m.add_lwpolyline([(0, 700), (10000, 700)], dxfattribs={"layer": "E-TRAY-700"})
        for x in range(0, 10001, 300): m.add_line((x, 0), (x, 700), dxfattribs={"layer": "E-TRAY-700"})
    else:
        m.add_lwpolyline([(0, 0), (6000, 0), (6000, 4000)], dxfattribs={"layer": "E-TRAY-700"})            # 10 m
        if dup: m.add_lwpolyline([(0, 0), (6000, 0), (6000, 4000)], dxfattribs={"layer": "E-TRAY-700"})   # ganda
    m.add_line((0, 100), (5000, 100), dxfattribs={"layer": "E-TRAY-200"})
    b = d.blocks.new("TRAYBLK"); b.add_line((0, 0), (1000, 0), dxfattribs={"layer": "0"})
    m.add_blockref("TRAYBLK", (0, 500), dxfattribs={"layer": "E-TRAY-200"})                                 # +1 m lewat blok
    m.add_line((0, 0, 0), (3000, 4000, 0), dxfattribs={"layer": "E-CONDUIT"})                               # 5 m horizontal
    if z: m.add_line((0, 0, 0), (0, 0, 2500), dxfattribs={"layer": "E-CONDUIT"})                            # 2.5 m vertikal
    for n, q in (("SUH-42GC", 3), ("SUH-46BC", 2)):
        d.blocks.new(n)
        for i in range(q): m.add_blockref(n, (i * 100, 900))
    d.saveas(path); return path

def cfg(): return T.load_config()

class CAD(unittest.TestCase):
    def setUp(self): self.d = tempfile.mkdtemp()
    def test_ukur_dedup_blok_3d_support(self):
        rec, sup, leg = cad.measure(dxf(os.path.join(self.d, "a.dxf")), cfg())
        self.assertEqual(rec["tray_m"], {"700": 10.0, "200": 6.0})            # garis ganda dibuang, blok ikut
        self.assertEqual((rec["conduit_h_m"], rec["conduit_v_m"]), (5.0, 2.5))
        self.assertEqual({k: v["qty"] for k, v in sup.items()}, {"SUH-42GC": 3, "SUH-46BC": 2})
        self.assertTrue(any("garis ganda" in f for f in rec["flags"]))
    def test_ladder_rel_dibagi_dua(self):
        rec, _, _ = cad.measure(dxf(os.path.join(self.d, "b.dxf"), ladder=True), cfg())
        self.assertAlmostEqual(rec["tray_m"]["700"], 10.0, delta=0.2); self.assertTrue(any("ladder" in f for f in rec["flags"]))
    def test_satuan_kosong_dan_2d(self):
        rec, _, _ = cad.measure(dxf(os.path.join(self.d, "c.dxf"), units=0, z=False), cfg())
        self.assertTrue(any("INSUNITS" in f for f in rec["flags"])); self.assertTrue(any("2D" in f for f in rec["flags"])); self.assertEqual(rec["conduit_v_m"], 0)
    def test_satuan_meter(self):
        rec, _, _ = cad.measure(dxf(os.path.join(self.d, "e.dxf"), units=6), cfg())      # 6 = meter: nilai gambar dianggap meter
        self.assertGreater(rec["tray_m"]["700"], 1000)
    def test_dedup_terpisah_per_role_dan_toleransi_meter(self):
        path = os.path.join(self.d, "meter.dxf")
        doc = ezdxf.new("R2018", setup=True); doc.units = 6
        doc.layers.add("E-TRAY-700"); doc.layers.add("E-CONDUIT")
        m = doc.modelspace()
        m.add_line((0, 0), (10, 0), dxfattribs={"layer": "E-TRAY-700"})
        m.add_line((0, 0), (10, 0), dxfattribs={"layer": "E-CONDUIT"})
        m.add_line((0.0004, 0), (10.0004, 0), dxfattribs={"layer": "E-TRAY-700"})
        doc.saveas(path)
        rec, _, _ = cad.measure(path, cfg())
        self.assertEqual(rec["tray_m"]["700"], 10.0)
        self.assertEqual(rec["conduit_h_m"], 10.0)
        self.assertTrue(any("garis ganda" in f for f in rec["flags"]))
    def test_tumpang_tindih_sebagian_dihitung_satu_kali(self):
        path = os.path.join(self.d, "partial-overlap.dxf")
        doc = ezdxf.new("R2018", setup=True); doc.units = 6
        doc.layers.add("E-TRAY-700")
        m = doc.modelspace()
        m.add_line((0, 0), (6, 0), dxfattribs={"layer": "E-TRAY-700"})
        m.add_line((4, 0), (10, 0), dxfattribs={"layer": "E-TRAY-700"})
        doc.saveas(path)
        rec, _, _ = cad.measure(path, cfg())
        self.assertEqual(rec["tray_m"]["700"], 10.0)
        self.assertTrue(any("tumpang tindih" in flag for flag in rec["flags"]))
    def test_include_window_memotong_di_batas_jendela(self):
        path = os.path.join(self.d, "window.dxf")
        doc = ezdxf.new("R2018", setup=True); doc.units = 6
        doc.layers.add("E-TRAY-700")
        doc.modelspace().add_line((0, 0), (10, 0), dxfattribs={"layer": "E-TRAY-700"})
        doc.saveas(path)
        config = cfg(); config["include_windows"] = [[2, -1, 8, 1]]
        rec, _, _ = cad.measure(path, config)
        self.assertEqual(rec["tray_m"]["700"], 6.0)
        self.assertEqual(rec["runs"][0]["bbox_du"], [2.0, 0.0, 8.0, 0.0])
    def test_entitas_tidak_didukung_diberi_peringatan(self):
        path = os.path.join(self.d, "unsupported.dxf")
        doc = ezdxf.new("R2018", setup=True); doc.layers.add("E-TRAY-700")
        doc.modelspace().add_text("label", dxfattribs={"layer": "E-TRAY-700"})
        doc.saveas(path)
        rec, _, _ = cad.measure(path, cfg())
        self.assertTrue(any("entitas TEXT" in f and "tidak lengkap" in f for f in rec["flags"]))
    def test_gagal_flattening_dilaporkan(self):
        from unittest.mock import patch
        entity = ezdxf.new("R2018").modelspace().add_arc((0, 0), 1, 0, 90)
        with patch.object(cad, "_flat_len", side_effect=RuntimeError("broken geometry")):
            with self.assertRaises(cad.CadGeometryError):
                cad.segments(entity)
    def test_scan_menemukan_layer_dan_support(self):
        s = cad.scan(dxf(os.path.join(self.d, "f.dxf")), cfg()); roles = {l["layer"]: l["role_guess"] for l in s["layers"]}
        self.assertEqual((roles["E-TRAY-700"], roles["E-CONDUIT"]), ("tray", "conduit")); self.assertEqual(s["support_blocks"]["SUH-42GC"], 3); self.assertTrue(s["has_3d"])
        tray_layer = next(layer for layer in s["layers"] if layer["layer"] == "E-TRAY-700")
        self.assertEqual(tray_layer["bbox_du"], [0.0, 0.0, 6000.0, 4000.0])
    def test_crosscheck_spacing(self):
        d = dxf(os.path.join(self.d, "g.dxf")); out = os.path.join(self.d, "o"); os.makedirs(out)
        r = T.plan_cad(d, cfg(), out); self.assertTrue(any("700W" in c and "PERINGATAN" in c for c in r["checks"]))   # 10 m / 3 = 3.33 m
    def test_dwg_gagal_memberi_instruksi(self):
        bad = os.path.join(self.d, "x.dwg"); open(bad, "wb").write(b"AC1032" + os.urandom(5000))
        with self.assertRaises(cad.ConvertError) as cm: cad.convert_dwg(bad, os.path.join(self.d, "o"))
        self.assertIn("DXF", str(cm.exception))
    def test_cli_plan_execute_dxf(self):
        f = dxf(os.path.join(self.d, "h.dxf")); s = os.path.join(os.path.dirname(__file__), "..", "scripts", "takeoff.py")
        subprocess.run([sys.executable, s, "plan", f, "--out", self.d], check=True, capture_output=True)
        self.assertNotEqual(subprocess.run([sys.executable, s, "execute", f, "--out", self.d], capture_output=True).returncode, 0)
        subprocess.run([sys.executable, s, "execute", f, "--out", self.d, "--confirm"], check=True, capture_output=True)
        self.assertTrue(os.path.exists(os.path.join(self.d, "Takeoff_Tray_Conduit.xlsx")))
        from openpyxl import load_workbook
        workbook = load_workbook(os.path.join(self.d, "Takeoff_Tray_Conduit.xlsx"), read_only=True)
        self.assertIn("Runs Tray", workbook.sheetnames)
        self.assertIn("Tray per Orientasi", workbook.sheetnames)
        workbook.close()
    def test_cli_execute_menolak_dxf_berbeda(self):
        f1 = dxf(os.path.join(self.d, "first.dxf")); f2 = dxf(os.path.join(self.d, "second.dxf"), dup=False)
        s = os.path.join(os.path.dirname(__file__), "..", "scripts", "takeoff.py")
        subprocess.run([sys.executable, s, "plan", f1, "--out", self.d], check=True, capture_output=True)
        r = subprocess.run([sys.executable, s, "execute", f2, "--out", self.d, "--confirm"], capture_output=True, text=True)
        self.assertNotEqual(r.returncode, 0); self.assertIn("tidak cocok", r.stderr)
        self.assertFalse(os.path.exists(os.path.join(self.d, "Takeoff_Tray_Conduit.xlsx")))

class Converter(unittest.TestCase):
    def setUp(self): self.d = tempfile.mkdtemp()
    def test_find_converter_env(self):
        from unittest import mock
        fake = os.path.join(self.d, "dwg2dxf"); open(fake, "w").close()
        with mock.patch.dict(os.environ, {"DWG2DXF_BIN": fake}):
            self.assertEqual(cad.find_converter(), ("dwg2dxf", fake))
    def test_convert_dwg_runs_converter_and_reads_back_dxf(self):
        from types import SimpleNamespace
        from unittest import mock
        source = os.path.join(self.d, "drawing.dwg")
        with open(source, "wb") as stream:
            stream.write(b"AC1032" + b"0" * 2000)
        converter = os.path.join(self.d, "dwg2dxf")
        open(converter, "w").close()

        def fake_run(args, **kwargs):
            output = args[args.index("-o") + 1]
            ezdxf.new("R2018", setup=True).saveas(output)
            return SimpleNamespace(stderr="", stdout="")

        with mock.patch.dict(os.environ, {"DWG2DXF_BIN": converter}), \
                mock.patch.object(cad.subprocess, "run", side_effect=fake_run):
            result = cad.convert_dwg(source, os.path.join(self.d, "out"))
        self.assertTrue(os.path.exists(result))
        self.assertEqual(ezdxf.readfile(result).dxfversion, "AC1032")
    def test_find_converter_kosong(self):
        from unittest import mock
        env = {k: v for k, v in os.environ.items() if k not in ("DWG2DXF_BIN", "ODA_BIN")}
        with mock.patch.dict(os.environ, env, clear=True), \
                mock.patch.object(cad.shutil, "which", return_value=None), \
                mock.patch("convert_local.find_oda", return_value=None):
            self.assertEqual(cad.find_converter(), (None, None))
    def test_pesan_error_menyebut_setup_dan_konversi_lokal(self):
        from unittest import mock
        bad = os.path.join(self.d, "x.dwg"); open(bad, "wb").write(b"AC1032" + os.urandom(2000))
        with mock.patch.object(cad, "find_converter", return_value=(None, None)):
            with self.assertRaises(cad.ConvertError) as cm: cad.convert_dwg(bad, os.path.join(self.d, "o"))
        self.assertIn("setup_converter.sh", str(cm.exception)); self.assertIn("convert_local.py", str(cm.exception))
    def test_convert_local_find_oda_env(self):
        from unittest import mock
        import convert_local as CL
        fake = os.path.join(self.d, "ODAFileConverter"); open(fake, "w").close()
        with mock.patch.dict(os.environ, {"ODA_BIN": fake}): self.assertEqual(CL.find_oda(), fake)
    def test_convert_local_verify_menolak_file_kecil(self):
        import convert_local as CL
        small = os.path.join(self.d, "a.dxf"); open(small, "w").write("0")
        self.assertFalse(CL.verify(small)[0]); self.assertFalse(CL.verify(os.path.join(self.d, "none.dxf"))[0])

if __name__ == "__main__": unittest.main()

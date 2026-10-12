"""Tes conduit per ukuran x jenis. Jalankan: python3 -m unittest discover -s tests"""
import json, os, shutil, subprocess, sys, tempfile, unittest
sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "scripts"))
import conduit as C

def W(t, x, y, h=35, c=90.0): return dict(t=t, x=x, y=y, w=len(t) * 20, h=h, c=c)

class Parse(unittest.TestCase):
    def test_pasangan_ukuran_jenis_dan_abaikan_elevasi(self):
        words = [W("42mm\u00d8", 100, 100), W("SIGNAL", 100, 150), W("EL-12838mm", 100, 200),
                 W("104mmo", 500, 100), W("POWER", 502, 152), W("28mm\u00d8", 900, 100), W("CONTROL", 905, 148)]
        lab, unp = C.parse_labels(words)
        self.assertEqual(sorted((l["size"], l["type"]) for l in lab), [(28, "CONTROL"), (42, "SIGNAL"), (104, "POWER")])
        self.assertEqual(unp, [])
    def test_ocr_salah_huruf_dan_ukuran_tanpa_jenis(self):
        lab, unp = C.parse_labels([W("7Omm\u00d8", 100, 100), W("POWFR", 100, 150), W("22mm", 800, 100)])
        self.assertEqual((lab[0]["size"], lab[0]["type"]), (70, "POWER"))
        self.assertEqual([u["size"] for u in unp], [22])                       # tidak dibuang diam-diam
    def test_ukuran_tak_umum_dan_conf_rendah_ditandai(self):
        lab, _ = C.parse_labels([W("700mm", 100, 100, c=10), W("POWER", 100, 150)])
        self.assertFalse(lab[0]["known"]); self.assertTrue(lab[0]["low_conf"])
    def test_jenis_dari_baris_lain_tidak_dipakai_dua_kali(self):
        lab, unp = C.parse_labels([W("28mm", 100, 100), W("42mm", 100, 130), W("SIGNAL", 100, 170)])
        self.assertEqual(len(lab), 1); self.assertEqual(len(unp), 1)

class Warna(unittest.TestCase):
    def test_variasi_warna_gambar_proyek_terbaca(self):
        import takeoff as T
        cfg = T.load_config()
        for c in ((0, 1, 0), (0, 1, 0.1), (0, 1, 0.13), (0.02, 1, 0)): self.assertTrue(T.iscol(c, cfg["colors"]["conduit"], cfg["color_tol"]), c)
        for c in ((1, 0, 1), (1, 0, 0.97), (1, 0.14, 1)): self.assertTrue(T.iscol(c, cfg["colors"]["tray"], cfg["color_tol"]), c)
        self.assertFalse(T.iscol((1, 0, 0), cfg["colors"]["tray"], cfg["color_tol"]))      # merah murni bukan tray
        self.assertFalse(T.iscol((0, 0, 1), cfg["colors"]["conduit"], cfg["color_tol"]))

class Alokasi(unittest.TestCase):
    def test_proporsional_dan_jumlah_tetap(self):
        r = C.allocate({(22, "SIGNAL"): 3, (42, "POWER"): 1}, 40.0)
        self.assertAlmostEqual(r[(22, "SIGNAL")], 30.0); self.assertAlmostEqual(sum(r.values()), 40.0)
    def test_tanpa_label_tidak_dialokasi(self):
        self.assertEqual(C.allocate({}, 10.0), {})

def fake_plan(d, dup=True):
    row = lambda s, role, items=None: {"row": {"sheet": s, "role": role, "items": items or {"conduit": 5}},
                                       "rec": {"role": role, "conduit_h_m": 100.0 if role == "plan" else 0, "conduit_v_m": 20.0 if role == "section" else 0}}
    with open(os.path.join(d, "pages.jsonl"), "w") as f:
        for r in (row(1, "plan"), row(2, "section"), row(3, "plan")): f.write(json.dumps(r) + "\n")
    sets = [{"set": 1, "pages": [1, 2], "name": "MBR TANKS (SOUTH)", "include": True},
            {"set": 2, "pages": [3, 3], "name": "MBR TANKS (SOUTH)", "include": False}]
    with open(os.path.join(d, "plan.json"), "w") as f: json.dump({"sets": sets}, f)

class Agregat(unittest.TestCase):
    def setUp(self): self.d = tempfile.mkdtemp(); fake_plan(self.d)
    def labels(self):
        L = lambda s, t, conf=90: {"size": s, "type": t, "x": 0, "y": 0, "conf": conf, "known": s in C.KNOWN_SIZES, "low_conf": conf < C.LOW_CONF}
        return {1: {"sheet": 1, "labels": [L(22, "SIGNAL"), L(22, "SIGNAL"), L(42, "POWER")], "unpaired": []},
                2: {"sheet": 2, "labels": [L(28, "CONTROL")], "unpaired": []},
                3: {"sheet": 3, "labels": [], "unpaired": []}}
    def test_per_set(self):
        rows = C.aggregate(self.d, self.labels()); s1, s2 = rows
        self.assertAlmostEqual(s1["h"][(22, "SIGNAL")], 100 * 2 / 3); self.assertAlmostEqual(s1["v"][(28, "CONTROL")], 20.0)
        self.assertEqual(s1["count"][(22, "SIGNAL")], 2)
        self.assertEqual(s2["h_none"], 100.0); self.assertTrue(any("tanpa label" in f for f in s2["flags"]))
    def test_halaman_belum_ocr_dilaporkan(self):
        lab = self.labels(); del lab[2]
        self.assertTrue(any("belum di-OCR" in f for f in C.aggregate(self.d, lab)[0]["flags"]))
    def test_excel_rumus_tanpa_error_dan_set_duplikat_tidak_dihitung(self):
        import openpyxl
        rows = C.aggregate(self.d, self.labels()); out = tempfile.mkdtemp()
        path = C.write_excel(out, rows, lambda n: "MBR South")
        wb = openpyxl.load_workbook(path); self.assertEqual(wb.sheetnames, ["Rekap Conduit per Area", "Detail Set"])
        ws = wb["Detail Set"]; self.assertEqual([r[4].value for r in ws.iter_rows(min_row=2)][:6], ["Y"] * 3 + ["N"] * 3)
        recalc = "/mnt/skills/public/xlsx/scripts/recalc.py"
        if os.path.exists(recalc):
            r = subprocess.run([sys.executable, recalc, path, "60"], capture_output=True, text=True); self.assertIn('"total_errors": 0', r.stdout)
            ws = openpyxl.load_workbook(path, data_only=True)["Rekap Conduit per Area"]
            vals = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, (int, float))]
            self.assertIn(4, vals); self.assertIn(100, vals); self.assertIn(20, vals)   # set 2 (N) tidak ikut: horizontal 100 bukan 200

class Gerbang(unittest.TestCase):
    def test_execute_tanpa_confirm_ditolak(self):
        d = tempfile.mkdtemp(); fake_plan(d); s = os.path.join(os.path.dirname(__file__), "..", "scripts", "conduit.py")
        r = subprocess.run([sys.executable, s, "execute", "x.pdf", "--plan-dir", d, "--out", d], capture_output=True, text=True)
        self.assertNotEqual(r.returncode, 0); self.assertIn("confirm", r.stderr + r.stdout)

@unittest.skipUnless(shutil.which("tesseract"), "tesseract tidak terpasang")
class OCRNyata(unittest.TestCase):
    def test_label_pada_pdf_sintetis(self):
        import pymupdf as fitz
        doc = fitz.open(); p = doc.new_page(width=1191, height=842)
        for i, (sz, tp) in enumerate((("42", "SIGNAL"), ("104", "POWER"))):
            p.insert_text((100 + i * 300, 200), f"{sz}mm", fontsize=14); p.insert_text((100 + i * 300, 222), tp, fontsize=14)
        lab, _ = C.parse_labels(C.ocr_words(p))
        self.assertEqual(sorted((l["size"], l["type"]) for l in lab), [(42, "SIGNAL"), (104, "POWER")])

if __name__ == "__main__": unittest.main()

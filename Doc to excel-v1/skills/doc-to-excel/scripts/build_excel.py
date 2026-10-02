#!/usr/bin/env python3
"""records.json + schema.json -> Excel (validasi, sel ragu ditandai, sheet 'Perlu Verifikasi' dan 'Log').

  python3 build_excel.py --schema schema.json --records records.json --out /mnt/user-data/outputs/hasil.xlsx
      [--extract-json /tmp/doc2x/a.json ...]   # untuk Log metode/halaman/peringatan
      [--control jumlah=118 ...]                # total yang tertulis di dokumen, dibandingkan dengan jumlah baris/penjumlahan
      [--append-to template.xlsx --sheet "Rekap" --header-row 3 [--start-row 20]]   # tambah ke Excel yang sudah ada (salinan)

records.json = list objek; kunci = key kolom di schema. Kunci khusus (opsional):
  _source (nama file), _page, _conf (0-100 keyakinan OCR baris), _flags {key: "alasan ragu"}, _note.
Nilai yang tidak bisa dibaca: kosongkan (null) dan beri _flags; jangan ditebak. Skema: lihat references/schema.md."""
import argparse, datetime, difflib, json, os, re, shutil, sys
from copy import copy
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

FONT = "Arial"
YELLOW = PatternFill("solid", fgColor="FFF2CC")
HEAD = PatternFill("solid", fgColor="305496")
MONTHS = {"jan": 1, "feb": 2, "mar": 3, "apr": 4, "mei": 5, "may": 5, "jun": 6, "jul": 7, "agu": 8, "ags": 8, "aug": 8,
          "sep": 9, "okt": 10, "oct": 10, "nov": 11, "des": 12, "dec": 12}


def norm(s):
    return re.sub(r"[^a-z0-9]", "", str(s or "").lower())


def empty(v):
    return v is None or (isinstance(v, str) and not v.strip())


def parse_number(v, locale="id"):
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return v
    t = str(v).strip()
    neg = (t.startswith("(") and t.endswith(")")) or t.startswith("-") or t.startswith("−")
    t = re.sub(r"(?i)\b(rp|idr|usd)\b\.?|[$€\s]", "", t)           # simbol mata uang dan spasi boleh
    if re.search(r"[^\d.,()\-−+]", t) or not re.search(r"\d", t):   # huruf di dalam angka = kemungkinan salah baca OCR (1O, 2S)
        raise ValueError(v)
    t = re.sub(r"[^\d.,]", "", t)
    if locale == "en":
        t = t.replace(",", "")
    elif locale == "id":
        t = t.replace(".", "").replace(",", ".")
    else:                                            # auto
        if "," in t and "." in t:
            t = t.replace(".", "").replace(",", ".") if t.rfind(",") > t.rfind(".") else t.replace(",", "")
        elif t.count(",") + t.count(".") == 1:
            sep = "," if "," in t else "."
            t = t.replace(sep, "") if len(t.split(sep)[1]) == 3 else t.replace(",", ".")   # 1.234 = ribuan, 12,5 = desimal
        else:
            t = t.replace(",", "").replace(".", "")
    n = float(t)
    n = -n if neg else n
    return int(n) if n == int(n) else n


def parse_date(v):
    if isinstance(v, datetime.datetime):
        return v.date()
    if isinstance(v, datetime.date):
        return v
    t = str(v).strip()
    m = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})", t)
    if m:
        return datetime.date(int(m[1]), int(m[2]), int(m[3]))
    m = re.match(r"^(\d{1,2})[/\-.](\d{1,2})[/\-.](\d{2,4})$", t)          # hari-bulan-tahun (kebiasaan Indonesia)
    if m:
        y = int(m[3]) + (2000 if int(m[3]) < 100 else 0)
        return datetime.date(y, int(m[2]), int(m[1]))
    m = re.match(r"^(\d{1,2})\s+([A-Za-z]{3,9})\.?,?\s+(\d{2,4})$", t)
    if m and MONTHS.get(m[2].lower()[:3]):
        y = int(m[3]) + (2000 if int(m[3]) < 100 else 0)
        return datetime.date(y, MONTHS[m[2].lower()[:3]], int(m[1]))
    raise ValueError(v)


def validate(records, schema, thr):
    """Kembalikan (baris_bersih, flags[(idx, key, alasan)]). Tidak membuang baris, tidak menebak nilai."""
    cols, loc = schema["columns"], schema.get("number_locale", "id")
    rows, flags = [], []
    for i, rec in enumerate(records):
        out, rf = {}, {}
        for c in cols:
            k, t, raw = c["key"], c.get("type", "text"), rec.get(c["key"])
            if empty(raw):
                out[k] = None
                if c.get("required"):
                    rf[k] = "kolom wajib kosong"
                continue
            try:
                if t in ("number", "integer"):
                    val = parse_number(raw, loc)
                    if t == "integer" and val != int(val):
                        raise ValueError(raw)
                    val = int(val) if t == "integer" else val
                    if "min" in c and val < c["min"] or "max" in c and val > c["max"]:
                        rf[k] = f"di luar rentang ({c.get('min', '-')}..{c.get('max', '-')})"
                elif t == "date":
                    val = parse_date(raw)
                else:
                    val = re.sub(r"\s+", " ", str(raw)).strip()
                    if t == "enum":
                        opts = c.get("enum", [])
                        hit = next((o for o in opts if o.lower() == val.lower()), None)
                        if not hit:
                            near = difflib.get_close_matches(val, opts, n=1, cutoff=0.85)
                            if near:
                                rf[k] = f"dikoreksi ke '{near[0]}' dari '{val}' (ejaan mirip)"
                                hit = near[0]
                            else:
                                rf[k] = f"nilai '{val}' tidak ada di daftar ({', '.join(opts[:6])}...)"
                        val = hit or val
                    if c.get("pattern") and not re.fullmatch(c["pattern"], val):
                        rf[k] = f"format tidak sesuai pola {c['pattern']}"
            except (ValueError, TypeError):
                val = str(raw)
                rf[k] = f"tidak terbaca sebagai {t}: '{raw}'"
            out[k] = val
        for k, why in (rec.get("_flags") or {}).items():
            rf[k] = (rf[k] + "; " if k in rf else "") + why
        cf = rec.get("_conf")
        if cf is not None and cf < thr and not rf:
            rf[cols[0]["key"]] = f"keyakinan OCR baris rendah ({cf})"
        for k in ("_source", "_page", "_conf", "_note"):
            out[k] = rec.get(k)
        rows.append(out)
        flags += [(i, k, w) for k, w in rf.items()]
    dk = schema.get("dedupe_key")
    if dk:
        seen = {}
        for i, r in enumerate(rows):
            key = tuple(norm(r.get(k)) for k in dk)
            if all(key):
                if key in seen:
                    flags.append((i, dk[0], f"duplikat dengan baris data ke-{seen[key] + 1} (tidak dihapus, cek manual)"))
                else:
                    seen[key] = i
    return rows, flags


def style_cell(c, col):
    c.font = Font(name=FONT, size=10)
    t = col.get("type", "text")
    if col.get("format"):
        c.number_format = col["format"]
    elif t == "integer":
        c.number_format = "#,##0"
    elif t == "number":
        c.number_format = "#,##0.00"
    elif t == "date":
        c.number_format = "dd-mmm-yyyy"


def put(ws, r, cidx, val, col, flagtxt):
    c = ws.cell(row=r, column=cidx, value=val)
    style_cell(c, col)
    if flagtxt:
        c.fill = YELLOW
        c.comment = Comment(flagtxt, "doc-to-excel")
    return c


def write_new(wb, schema, rows, flags, meta):
    ws = wb.active
    ws.title = schema.get("sheet", "Data")
    cols = schema["columns"]
    heads = [c.get("header", c["key"]) for c in cols] + (["Sumber", "Status"] if meta else [])
    ws.append(heads)
    for c in ws[1]:
        c.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        c.fill = HEAD
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    fl = {}
    for i, k, w in flags:
        fl.setdefault((i, k), []).append(w)
    for i, r in enumerate(rows):
        for j, c in enumerate(cols, 1):
            put(ws, i + 2, j, r[c["key"]], c, "; ".join(fl.get((i, c["key"]), [])))
        if meta:
            src = f"{r.get('_source') or ''}" + (f" hlm {r['_page']}" if r.get("_page") else "")
            ws.cell(row=i + 2, column=len(cols) + 1, value=src.strip()).font = Font(name=FONT, size=9, color="808080")
            bad = any(ii == i for ii, _, _ in flags)
            sc = ws.cell(row=i + 2, column=len(cols) + 2, value="CEK" if bad else "OK")
            sc.font = Font(name=FONT, size=10, bold=bad, color="C00000" if bad else "2E7D32")
    last = len(rows) + 1
    tot = schema.get("totals") or []
    if tot and rows:
        ws.cell(row=last + 1, column=1, value="Total").font = Font(name=FONT, bold=True)
        for j, c in enumerate(cols, 1):
            if c["key"] in tot:
                cell = ws.cell(row=last + 1, column=j, value=f"=SUM({L(j)}2:{L(j)}{last})")
                style_cell(cell, c)
                cell.font = Font(name=FONT, bold=True)
                cell.border = Border(top=Side(style="thin"))
    for j, c in enumerate(cols, 1):
        ws.column_dimensions[L(j)].width = c.get("width", max(12, min(40, len(heads[j - 1]) + 4)))
    if meta:
        ws.column_dimensions[L(len(cols) + 1)].width = 28
        ws.column_dimensions[L(len(cols) + 2)].width = 9
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{L(len(heads))}{max(last, 2)}"
    return ws


def append_existing(path, out, schema, rows, flags, sheet, header_row, start_row):
    keep_vba = out.lower().endswith(".xlsm")
    shutil.copy(path, out)
    wb = load_workbook(out, keep_vba=keep_vba)
    ws = wb[sheet] if sheet else wb.active
    hdr = {norm(ws.cell(row=header_row, column=c).value): c for c in range(1, ws.max_column + 1)
           if ws.cell(row=header_row, column=c).value}
    cmap, missing = {}, []
    for c in schema["columns"]:
        col = hdr.get(norm(c.get("header", c["key"]))) or hdr.get(norm(c["key"]))
        (cmap.__setitem__(c["key"], col) if col else missing.append(c.get("header", c["key"])))
    if not cmap:
        sys.exit(f"Tidak ada header cocok di baris {header_row} sheet '{ws.title}'. Header terbaca: {list(hdr)[:10]}")
    r0 = start_row
    if not r0:
        r0 = header_row + 1
        while any(not empty(ws.cell(row=r0, column=cc).value) for cc in cmap.values()):
            first = str(ws.cell(row=r0, column=min(cmap.values())).value or "").strip().lower()
            if first.startswith(("total", "jumlah", "subtotal")):
                sys.exit(f"Baris Total di baris {r0}: menambah data di sini akan menimpanya. Tentukan --start-row (baris kosong) atau sisipkan baris dulu.")
            r0 += 1
    model = r0 - 1 if r0 - 1 > header_row else None
    fl = {}
    for i, k, w in flags:
        fl.setdefault((i, k), []).append(w)
    for i, r in enumerate(rows):
        for c in schema["columns"]:
            cc = cmap.get(c["key"])
            if not cc:
                continue
            cell = put(ws, r0 + i, cc, r[c["key"]], c, "; ".join(fl.get((i, c["key"]), [])))
            if model:                                   # ikuti gaya baris data sebelumnya
                src = ws.cell(row=model, column=cc)
                cell.font, cell.border, cell.alignment = copy(src.font), copy(src.border), copy(src.alignment)
                if not c.get("format") and src.number_format != "General":
                    cell.number_format = src.number_format
    return wb, ws, cmap, missing, r0


def add_review(wb, rows, flags, schema):
    if not flags:
        return
    ws = wb.create_sheet("Perlu Verifikasi")
    ws.append(["Baris data", "Kolom", "Nilai di Excel", "Alasan", "Sumber"])
    heads = {c["key"]: c.get("header", c["key"]) for c in schema["columns"]}
    for i, k, w in flags:
        r = rows[i]
        src = f"{r.get('_source') or ''}" + (f" hlm {r['_page']}" if r.get("_page") else "")
        ws.append([i + 1, heads.get(k, k), "" if r.get(k) is None else str(r.get(k)), w, src.strip()])
    for c in ws[1]:
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = HEAD
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name=FONT, size=10)
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for col, w in zip("ABCDE", (11, 22, 26, 70, 34)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"


def add_log(wb, rows, flags, schema, extracts, controls, notes):
    ws = wb.create_sheet("Log")
    ws.column_dimensions["A"].width = 150
    lines = [f"Log input data - {datetime.datetime.now():%Y-%m-%d %H:%M}",
             f"Baris ditulis: {len(rows)} | sel/baris bertanda: {len(flags)} | baris bermasalah: {len({i for i, _, _ in flags})}", ""]
    for e in extracts:
        d = json.load(open(e, encoding="utf-8"))
        meth = {}
        for p in d["pages"]:
            meth[p["method"]] = meth.get(p["method"], 0) + 1
        lines.append(f"Sumber {d['file']}: {len(d['pages'])} hlm, metode " + ", ".join(f"{k} x{v}" for k, v in meth.items()))
        lines += ["  PERINGATAN: " + w for w in d.get("warnings", [])]
    for k, stated in controls.items():
        c = next((x for x in schema["columns"] if x["key"] == k), None)
        if c is None:
            lines.append(f"Kontrol {k}: kolom tidak ada di skema")
            continue
        got = sum(r[k] for r in rows if isinstance(r.get(k), (int, float))) if c.get("type") in ("number", "integer") else len(rows)
        lines.append(f"Kontrol {k}: dokumen menulis {stated}, hasil input {got} (selisih {got - stated:+g})" + ("" if got == stated else "  <- CEK"))
    lines += notes
    for ln in lines:
        ws.append([ln])
    ws["A1"].font = Font(name=FONT, bold=True, size=12)


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--schema", required=True)
    ap.add_argument("--records", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--extract-json", nargs="*", default=[])
    ap.add_argument("--control", action="append", default=[], help="kolom=nilai total yang tertulis di dokumen")
    ap.add_argument("--conf-threshold", type=float, default=75)
    ap.add_argument("--append-to")
    ap.add_argument("--sheet")
    ap.add_argument("--header-row", type=int, default=1)
    ap.add_argument("--start-row", type=int)
    a = ap.parse_args()
    schema = json.load(open(a.schema, encoding="utf-8"))
    records = json.load(open(a.records, encoding="utf-8"))
    rows, flags = validate(records, schema, a.conf_threshold)
    controls = {k: float(v) for k, v in (x.split("=", 1) for x in a.control)}
    notes = []
    os.makedirs(os.path.dirname(os.path.abspath(a.out)), exist_ok=True)
    if a.append_to:
        wb, ws, cmap, missing, r0 = append_existing(a.append_to, a.out, schema, rows, flags, a.sheet, a.header_row, a.start_row)
        notes.append(f"Ditambahkan ke salinan {os.path.basename(a.append_to)} sheet '{ws.title}' mulai baris {r0} ({len(rows)} baris).")
        if missing:
            notes.append("Kolom skema tanpa header padanan di Excel (tidak ditulis): " + ", ".join(missing))
    else:
        wb = Workbook()
        write_new(wb, schema, rows, flags, not schema.get("no_meta"))
    add_review(wb, rows, flags, schema)
    add_log(wb, rows, flags, schema, a.extract_json, controls, notes)
    wb.save(a.out)
    nbad = len({i for i, _, _ in flags})
    print(f"{a.out}: {len(rows)} baris, {nbad} baris perlu verifikasi ({len(flags)} sel/catatan).")
    for ln in notes:
        print(ln)
    if schema.get("totals") and not a.append_to:
        print("Ada rumus Total: jalankan recalc.py dari skill xlsx pada file ini.")

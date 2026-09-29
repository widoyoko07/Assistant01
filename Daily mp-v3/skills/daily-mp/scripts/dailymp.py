#!/usr/bin/env python3
"""Daily MP - rekap manpower harian semua subcontractor (format Excel rekap manpower).

  python dailymp.py --dir /tmp/mp --date 2026-09-28
  python dailymp.py --dir /tmp/mp --date 2026-09-28 --detail
  python dailymp.py --dir /tmp/mp --date 2026-09-28 --xlsx out.xlsx
  python dailymp.py --dir /tmp/mp --summary                 # total per subcontractor per tanggal
  python dailymp.py --dir /tmp/mp --date 2026-09-28 --map "Driver=Staff" --map "Erector=Skilled Workers"

File di --dir harus diawali kode subcontractor: "SILOG TBM 20260928 Posting.txt", "TPE_20260928.txt",
"DJK MP Service 26-28 Sep 2026.txt" (DJK boleh multi-hari), dst.
"""
import re, os, json, argparse, difflib, glob
from datetime import date as _date

SUBS = ["SILOG", "TODJO", "DHJ", "WKP", "TPE", "WME", "DJK", "BCP", "BME"]
ALIAS = {"TOHOMA": "TODJO", "TOHOMA-DWIDJO": "TODJO"}
NEW = " (role baru)"

ROWS = [
    ("Indirect", "Project Manager"), ("Indirect", "Deputy Project Manager"),
    ("Indirect", "Construction Manager"), ("Indirect", "Deputy CM"),
    ("Indirect", "Site Engineer"), ("Indirect", "HSE Manager"), ("Indirect", "HSE Engineer"),
    ("Indirect", "Survey Manager"), ("Indirect", "Surveyor"), ("Indirect", "QS Department"),
    ("Indirect", "Engineering Department"), ("Indirect", "QA/QC Department"),
    ("Indirect", "Administration Department"), ("Indirect", "Supervisor"),
    ("Direct", "Foreman"), ("Direct", "Rebarman"), ("Direct", "Carpenter"), ("Direct", "Stone Masonry"),
    ("Direct", "Scaffolder"), ("Direct", "Skilled Workers"), ("Direct", "Common Labor"),
    ("", "Security"), ("Indirect", "Staff"), ("Direct", "Heavy Equipment Operator"),
]
POS = [p for _, p in ROWS]
TYPE_OF = {p: t for t, p in ROWS}
SUM_ROWS = [("Indirect", "Indirect (tanpa rincian posisi)"), ("Direct", "Direct (tanpa rincian posisi)")]

ROLE_MAP = {
    "pm": "Project Manager", "project manager": "Project Manager",
    "dpm": "Deputy Project Manager", "deputy project manager": "Deputy Project Manager",
    "cm": "Construction Manager", "construction manager": "Construction Manager",
    "dcm": "Deputy CM", "deputy cm": "Deputy CM",
    # SM = Site Manager, sengaja dimasukkan ke Site Engineer (kebijakan user)
    "sm": "Site Engineer", "site manager": "Site Engineer", "se": "Site Engineer",
    "site engineer": "Site Engineer", "engineer": "Site Engineer", "enginer": "Site Engineer",
    "eng": "Site Engineer", "rigging eng": "Site Engineer", "pc": "Site Engineer",
    "ppc": "Site Engineer", "project control": "Site Engineer",
    "hse manager": "HSE Manager",
    "survey manager": "Survey Manager", "surveyor": "Surveyor", "team surveyor": "Surveyor",
    "ass surveyor": "Surveyor", "qs": "QS Department",
    "qc": "QA/QC Department", "qa": "QA/QC Department", "qc lead": "QA/QC Department",
    "admin": "Administration Department", "docon": "Administration Department",
    "matecon": "Staff", "matkon": "Staff", "material control": "Staff",
    "spv": "Supervisor", "supervisor": "Supervisor",
    "foreman": "Foreman", "foremen": "Foreman", "formen": "Foreman", "fm": "Foreman",
    "si": "Foreman",  # SI = Subcontractor Independen
    "rebarman": "Rebarman", "carpenter": "Carpenter", "stone masonry": "Stone Masonry", "mason": "Stone Masonry",
    "scaffolder": "Scaffolder", "scf": "Scaffolder", "insp scaffolding": "Scaffolder",
    "inspector scaffolding": "Scaffolder", "inspektor scaffolder": "Scaffolder", "inspector scf": "Scaffolder",
    "teknisi": "Skilled Workers", "technician": "Skilled Workers", "rigger": "Skilled Workers",
    "welder": "Skilled Workers", "skill": "Skilled Workers", "skil": "Skilled Workers",
    "helper": "Skilled Workers", "hlp": "Skilled Workers", "fitter": "Skilled Workers",
    "ft": "Skilled Workers", "mw": "Skilled Workers", "mep": "Skilled Workers",
    "f": "Skilled Workers", "h": "Skilled Workers",
    "security": "Security",
}
# seksi -> posisi (dipakai untuk orang tanpa role). "team qc" memaksa semua anggota jadi QA/QC.
SECTION_RULES = [("scaffold", "Scaffolder"), ("surveyor", "Surveyor"), ("warehouse", "Common Labor"),
                 ("mechanical fitter", "Skilled Workers"), ("welder", "Skilled Workers"),
                 ("rigger", "Skilled Workers"), ("electri", "Skilled Workers"), ("operator", "Heavy Equipment Operator")]
ABSENT_RE = re.compile(r"\b(absen|sakit|izin|cuti|alpa)\b", re.I)
OVERRIDES = {}

def load_extra():
    p = os.path.join(os.path.dirname(os.path.abspath(__file__)), "roles.json")
    if os.path.exists(p):
        try:
            for k, v in json.load(open(p, encoding="utf-8")).get("role_map", {}).items():
                OVERRIDES[norm(k)] = v
        except Exception:
            pass

def norm(s):
    return re.sub(r"\s+", " ", re.sub(r"[_.\-/]+", " ", (s or "").lower())).strip()

def clean(text):
    text = re.sub(r"[\u200b-\u200f\u2060\ufeff]", "", text).replace("\\", "")
    return re.sub(r"(?<=[^\s\d\n])(\d{1,3}\s*\.\s)", r"\n\1", text)   # item yang menempel jadi baris baru

def label_new(txt):
    txt = re.sub(r"\(.*?\)", "", txt).strip(" :*-")
    if not any(ch.isdigit() for ch in txt) and (txt.islower() or (txt.isupper() and len(txt) > 4)):
        txt = txt.title()
    return (txt or "Tanpa role") + NEW

def classify(role, section):
    r, s = norm(role), norm(section)
    if "team qc" in s:
        return "QA/QC Department"
    for key in (r, s):
        if key and key in OVERRIDES:
            return OVERRIDES[key]
    if r:
        if r.startswith(("helper", "hlp")) and "warehouse" in s: return "Common Labor"
        if r.startswith("5r"): return label_new(role)   # tim 5R (kebersihan), tanya user
        if r in ROLE_MAP: return ROLE_MAP[r]
        if r.startswith(("spv", "supervisor")): return "Supervisor"
        if "hse" in r.split() or r.startswith("hse"): return "HSE Manager" if "manager" in r else "HSE Engineer"
        if r.startswith("operator"): return "Heavy Equipment Operator"
        if r.startswith(("helper", "hlp")): return "Common Labor" if "warehouse" in s else "Skilled Workers"
        if r.startswith("fm"): return "Foreman"
        if "insp" in r and ("scaf" in r or "scf" in r): return "Scaffolder"
        if r.startswith(("qc", "qa")): return "QA/QC Department"
        return label_new(role)
    for sub, pos in SECTION_RULES:
        if sub in s: return pos
    if s in ROLE_MAP: return ROLE_MAP[s]
    return label_new(section)

# ---------------------------------------------------------------- parser
def find_stated(text):
    m = re.search(r"(?i)total\s*(?:all|manpower|mp|aktif)\s*[:=]?\s*(\d+)", text)
    return int(m.group(1)) if m else None

def mk(section, role, name):
    return {"section": section, "role": role, "name": name, "category": classify(role, section)}

def split_item(body):
    """'Role : Name' | 'Name ( Role )' | 'Name' -> (role, name)"""
    if ":" in body:
        r, n = body.split(":", 1); return (r.strip() or None), n.strip()
    m = re.match(r"^(.*?)\s*\(([^)]*)\)\s*$", body)
    if m: return (m.group(2).strip() or None), m.group(1).strip()
    return None, body.strip()

def parse_numbered(text):
    people, absent, section, off = [], [], "Manpower", False
    for raw in clean(text).splitlines():
        line = raw.strip()
        if not line: continue
        if re.match(r"(?i)^[-*•\s]*total\b", line): break
        if re.match(r"(?i)^(area\b|supporting team|plan\b|kondisi|cuaca|work description|weather|loc\.)", line): break
        item = re.match(r"^\d+\s*[.)]\s*(.*)$", line)
        if item:
            body = item.group(1).strip()
            if not body: continue
            if off or ABSENT_RE.search(body):
                absent.append(body); continue
            role, name = split_item(body)
            if name: people.append(mk(section, role, name))
            continue
        hdr = re.match(r"(?i)^(.+?)\s*:\s*\d+\s*person", line)
        if hdr:
            section, off = hdr.group(1).strip(), False; continue
        head = re.match(r"^[*•\-]\s*(.+)$", line)
        if head and not re.match(r"(?i)^crew", head.group(1)):
            section = head.group(1).strip(" :*"); off = norm(section).startswith("off")
    return people, absent

def parse_dhj(text):
    people, role, on = [], None, False
    for raw in clean(text).splitlines():
        line = raw.strip()
        if re.match(r"(?i)^man\s*power", line): on = True; continue
        if not on: continue
        if re.match(r"(?i)^(total|plan|kondisi)", line): break
        m = re.match(r"^([^=]*?)\s*=\s*(.*)$", line)
        if not m: continue
        if m.group(1).strip(): role = m.group(1).strip()
        if m.group(2).strip(): people.append(mk("Man Power", role, m.group(2).strip()))
    return people, []

def parse_todjo(text):
    people, absent, mode = [], [], None
    for raw in clean(text).splitlines():
        line = re.sub(r"^[.\s]+", "", raw.strip())
        if not line: continue
        if re.match(r"(?i)^total", line) or re.match(r"(?i)^work description", line): break
        if re.match(r"(?i)^man\s*power", line): mode = "staff"; continue
        if re.match(r"(?i)^crew", line): mode = "crew"; continue
        if re.match(r"(?i)^tidak\s+berangkat", line): mode = "absent"; continue
        if mode is None: continue
        if mode == "staff" and ":" in line:
            r, n = line.split(":", 1)
            if n.strip(): people.append(mk("Staff", r.strip(), n.strip()))
        elif mode in ("crew", "absent"):
            role, name = split_item(line)
            if not name: continue
            (absent if mode == "absent" else people).append(name if mode == "absent" else mk("Crew", role, name))
    return people, absent

def parse_summary(text):
    t = clean(text); out = {}
    for key, pat in (("indirect", r"indirect\s*[=:]\s*(\d+)"), ("direct", r"(?<!in)direct\s*[=:]\s*(\d+)"),
                     ("staff", r"staf+\s*[:=]\s*(\d+)"), ("worker", r"worker\s*[:=]\s*(\d+)"),
                     ("total", r"total\s*(?:mp|manpower)\s*[:=]?\s*(\d+)")):
        m = re.search(pat, t, re.I)
        if m: out[key] = int(m.group(1))
    ind = out.get("indirect", out.get("staff")); dr = out.get("direct", out.get("worker"))
    return {"indirect": ind, "direct": dr, "total": out.get("total")} if (ind is not None or dr is not None) else None

def parse_djk(text):
    days = {}
    for block in re.split(r"(?=TBM PAGI)", text):
        m = re.search(r"(?i)(senin|selasa|rabu|rabo|kamis|jumat|sabtu|minggu)\s+(\d{1,2})-(\d{1,2})-(\d{4})", block)
        if not m: continue
        d = f"{m.group(4)}-{int(m.group(3)):02d}-{int(m.group(2)):02d}"
        p, a = parse_numbered(block)
        days[d] = (p, a, find_stated(block))
    return days

# ---------------------------------------------------------------- dedupe
def nn(s): return re.sub(r"[^a-z]", "", (s or "").lower())
def nr(r):
    r = norm(r)
    return {"foremen": "foreman", "formen": "foreman", "fm": "foreman"}.get(r, r)

def dedupe(people):
    removed, kept, seen = [], [], set()
    for p in people:
        k = (nn(p["name"]), nr(p["role"]), p["category"])
        if k in seen: removed.append((p, "nama & posisi sama")); continue
        seen.add(k); kept.append(p)
    final = []
    for b in kept:
        hit = None
        for i, a in enumerate(final):
            if nn(a["name"]) != nn(b["name"]) and a["category"] == b["category"] and \
               difflib.SequenceMatcher(None, nn(a["name"]), nn(b["name"])).ratio() >= 0.85:
                hit = i; break
        if hit is None: final.append(b)
        else:
            a = final[hit]
            if a["role"] is None and b["role"]:
                final[hit] = b; removed.append((a, f"ejaan mirip dengan {b['name']}"))
            else: removed.append((b, f"ejaan mirip dengan {a['name']}"))
    warns, by = [], {}
    for p in final: by.setdefault(nn(p["name"]), []).append(p)
    for v in by.values():
        if len(v) > 1:
            warns.append("nama sama, posisi beda (tidak dihapus, cek manual): " +
                         "; ".join(f"{x['name']} - {x['role'] or x['section']}" for x in v))
    return final, removed, warns

# ---------------------------------------------------------------- load
def sub_of(fname):
    tok = re.split(r"[\s_\-]", os.path.basename(fname).upper())[0]
    tok = ALIAS.get(tok, tok)
    return tok if tok in SUBS else None

def load_dir(folder):
    res = {s: {} for s in SUBS}
    for f in sorted(glob.glob(os.path.join(folder, "*"))):
        sub = sub_of(f)
        if not sub: continue
        if f.lower().endswith(".pdf"):
            m = re.search(r"(20\d{2})(\d{2})(\d{2})", f)
            if m: res[sub].setdefault(f"{m.group(1)}-{m.group(2)}-{m.group(3)}", {"status": "pdf", "file": os.path.basename(f)})
            continue
        text = open(f, encoding="utf-8", errors="ignore").read()
        if sub == "DJK":
            for d, (p, a, st) in parse_djk(text).items():
                res[sub][d] = {"status": "ok", "people": p, "absent": a, "stated": st, "file": os.path.basename(f)}
            continue
        m = re.search(r"(20\d{2})(\d{2})(\d{2})", os.path.basename(f))
        if not m: continue
        d = f"{m.group(1)}-{m.group(2)}-{m.group(3)}"
        if sub in ("WME", "WKP"):
            s = parse_summary(text)
            res[sub][d] = {"status": "summary", "summary": s, "file": os.path.basename(f)} if s else \
                          {"status": "unknown", "file": os.path.basename(f)}
            continue
        fn = {"DHJ": parse_dhj, "TODJO": parse_todjo}.get(sub, parse_numbered)
        p, a = fn(text)
        res[sub][d] = ({"status": "ok", "people": p, "absent": a, "stated": find_stated(text), "file": os.path.basename(f)}
                       if p else {"status": "unknown", "file": os.path.basename(f)})
    return res

# ---------------------------------------------------------------- report
def build(res, d, detail):
    cols, notes, newroles = {}, [], {}
    for s in SUBS:
        r = res[s].get(d)
        if not r:
            cols[s] = ("missing", None); notes.append(f"**{s}**: belum ada data untuk {d}."); continue
        if r["status"] == "pdf":
            cols[s] = ("missing", None)
            notes.append(f"**{s}**: hanya ada PDF foto absensi ({r['file']}), tidak dibaca otomatis. Kirim Posting.txt agar bisa direkap."); continue
        if r["status"] == "unknown":
            cols[s] = ("missing", None); notes.append(f"**{s}**: format {r['file']} tidak dikenali."); continue
        if r["status"] == "summary":
            sm = r["summary"]; cols[s] = ("summary", sm)
            notes.append(f"**{s}**: laporan hanya berisi total (Indirect {sm['indirect']}, Direct {sm['direct']}"
                         f"{', Total ' + str(sm['total']) if sm['total'] else ''}), tanpa rincian posisi. Ditaruh di baris 'tanpa rincian posisi'.")
            continue
        people, removed, warns = dedupe(r["people"])
        counts = {}
        for p in people: counts[p["category"]] = counts.get(p["category"], 0) + 1
        cols[s] = ("ok", counts)
        n = [f"**{s}**: {len(people)} orang"]
        if r["stated"] is not None and r["stated"] != len(people):
            n.append(f"total di teks {r['stated']} (selisih {len(people) - r['stated']:+d})")
        if r["absent"]: n.append(f"{len(r['absent'])} tidak hadir dikeluarkan")
        if any(norm(p["role"]) == "sm" for p in people): n.append("SM (Site Manager) masuk Site Engineer")
        if any(norm(p["role"]) == "si" for p in people): n.append("SI (Subcontractor Independen) masuk Foreman")
        if removed:
            n.append(f"{len(removed)} double dihapus: " + "; ".join(f"{p['name']} [{w}]" for p, w in removed))
        n += warns
        notes.append(" | ".join(n))
        for p in people:
            if p["category"] not in POS:
                newroles.setdefault(p["category"], {}).setdefault(s, []).append(p["name"])
        if detail:
            det = {}
            for p in people: det.setdefault(p["category"], []).append(p["name"])
            notes.append(f"  Rincian {s}: " + "; ".join(f"{k} ({len(v)}): {', '.join(v)}" for k, v in det.items()))
    rows_def = list(ROWS)
    if any(k == "summary" for k, _ in cols.values()): rows_def += SUM_ROWS
    for nr_label in newroles: rows_def.append(("", nr_label))
    table = []
    for t, p in rows_def:
        vals = []
        for s in SUBS:
            kind, data = cols[s]
            if kind == "missing": vals.append("n/a")
            elif kind == "summary":
                vals.append(data["indirect"] if p == SUM_ROWS[0][1] and data["indirect"] is not None else
                            data["direct"] if p == SUM_ROWS[1][1] and data["direct"] is not None else "-")
            else: vals.append(data.get(p, 0))
        table.append((t, p, vals))
    return table, notes, newroles

def total(vals): return sum(v for v in vals if isinstance(v, int))
def col_totals(table): return [total([r[2][i] for r in table]) for i in range(len(SUBS))]

def md(table, d):
    head = "| Tipe | Posisi | " + " | ".join(SUBS) + " | Total |"
    out = [f"## Daily MP - {d}", "", head, "|---|---|" + "---:|" * (len(SUBS) + 1)]
    for t, p, v in table:
        out.append(f"| {t} | {p} | " + " | ".join(str(x) for x in v) + f" | {total(v)} |")
    ct = col_totals(table)
    out.append("| | **Total Manpower** | " + " | ".join(f"**{x}**" for x in ct) + f" | **{sum(ct)}** |")
    return "\n".join(out)

def tsv(table):
    lines = ["\t".join(["Tipe", "Posisi"] + SUBS + ["Total"])]
    for t, p, v in table: lines.append("\t".join([t, p] + [str(x) for x in v] + [str(total(v))]))
    ct = col_totals(table)
    lines.append("\t".join(["", "Total Manpower"] + [str(x) for x in ct] + [str(sum(ct))]))
    return "\n".join(lines)

def to_xlsx(path, table, d):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter as L
    wb = Workbook(); ws = wb.active; ws.title = "Daily MP"
    ws.append([f"Daily MP - {d}"]); ws["A1"].font = Font(bold=True, size=13)
    ws.append(["Tipe", "Posisi"] + SUBS + ["Total"])
    for c in ws[2]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="305496"); c.alignment = Alignment(horizontal="center")
    first = 3
    for i, (t, p, v) in enumerate(table):
        r = first + i
        ws.append([t, p] + v + [f"=SUM(C{r}:{L(2 + len(SUBS))}{r})"])
    last = first + len(table) - 1
    ws.append(["", "Total Manpower"] + [f"=SUM({L(3 + i)}{first}:{L(3 + i)}{last})" for i in range(len(SUBS) + 1)])
    for c in ws[ws.max_row]: c.font = Font(bold=True)
    ws.column_dimensions["A"].width = 11; ws.column_dimensions["B"].width = 34
    for i in range(len(SUBS) + 1): ws.column_dimensions[L(3 + i)].width = 9
    ws.freeze_panes = "C3"; wb.save(path)

if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--dir", required=True); ap.add_argument("--date")
    ap.add_argument("--detail", action="store_true"); ap.add_argument("--xlsx")
    ap.add_argument("--summary", action="store_true"); ap.add_argument("--map", action="append", default=[])
    a = ap.parse_args()
    load_extra()
    for m in a.map:
        k, v = m.split("=", 1)
        v = next((p for p in POS if p.lower() == v.strip().lower()), None)
        if not v: raise SystemExit(f"Posisi tidak dikenal di --map: {m}. Pilih salah satu: {', '.join(POS)}")
        OVERRIDES[norm(k)] = v
    res = load_dir(a.dir)
    dates = sorted({d for s in SUBS for d in res[s]})
    if a.summary:
        print("| Tanggal | " + " | ".join(SUBS) + " | Total |\n|---|" + "---:|" * (len(SUBS) + 1))
        for d in dates:
            table, _, _ = build(res, d, False); ct = col_totals(table)
            print(f"| {d} | " + " | ".join(str(x) if res[s].get(d) and res[s][d]["status"] in ("ok", "summary") else "n/a"
                                               for s, x in zip(SUBS, ct)) + f" | {sum(ct)} |")
        raise SystemExit
    d = a.date or (dates[-1] if dates else None)
    if not d: raise SystemExit("Tidak ada file data di folder.")
    table, notes, newroles = build(res, d, a.detail)
    print(md(table, d)); print("\n**Catatan:**"); [print("- " + n) for n in notes]
    if newroles:
        print("\n**PERLU KONFIRMASI - role baru (belum ada di aturan), sementara ditaruh di baris sendiri:**")
        for lab, subs in newroles.items():
            print(f"- {lab.replace(NEW, '')}: " + "; ".join(f"{s} ({len(n)}: {', '.join(n)})" for s, n in subs.items()))
        print("Tanyakan ke user masuk posisi mana untuk tiap role di atas. Jangan menebak.")
    print("\n### Blok TSV (copy, paste ke Excel)\n\n```\n" + tsv(table) + "\n```")
    if a.xlsx: to_xlsx(a.xlsx, table, d); print(f"\nExcel: {a.xlsx}")

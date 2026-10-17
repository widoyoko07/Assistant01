"""Jalankan: python3 -m unittest discover -s tests   (dari folder skill). Memakai file Excel sintetis."""
import json, os, subprocess, sys, tempfile, unittest, zipfile
import openpyxl
from openpyxl.styles import Font, PatternFill
SCRIPT = os.path.join(os.path.dirname(__file__), "..", "scripts", "combine_sheets.py")
sys.path.insert(0, os.path.dirname(SCRIPT))
import combine_sheets as C


def make(path, sheets, bold=False, color="FF0000"):
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    ws = wb.create_sheet("Data"); ws["A1"] = 7
    for name, val in sheets:
        w = wb.create_sheet(name); w["A1"] = val; w["A2"] = 3; w["A3"] = "=SUM(A1:A2)"; w["A4"] = "=Data!A1"; w["A5"] = "teks"
        w["A1"].font = Font(bold=bold, color=color); w["A2"].fill = PatternFill("solid", fgColor="FFFF00"); w.merge_cells("C1:D2")
        w.column_dimensions["B"].width = 25
    wb.save(path)


def run(*a): return subprocess.run([sys.executable, SCRIPT, *a], capture_output=True, text=True)


class Gabung(unittest.TestCase):
    def setUp(self):
        self.d = tempfile.mkdtemp(); self.out = os.path.join(self.d, "o")
        make(os.path.join(self.d, "1001 a.xlsx"), [("2 Jun", 10), ("1 Jun", 5)])
        make(os.path.join(self.d, "1002 b.xlsx"), [("3 Jun", 20), ("1 Jun", 6)], bold=True, color="0000FF")   # nama kembar + gaya beda
        open(os.path.join(self.d, "1003 rusak.xlsx"), "wb").write(b"bukan zip")

    def test_scan_urut_kembar_dan_masalah(self):
        r = C.scan(self.d, "jun")
        self.assertEqual([e["out_name"] for e in r["entries"]], ["1 Jun", "1 Jun (1002)", "2 Jun", "3 Jun"])
        self.assertEqual(len(r["problems"]), 1); self.assertTrue(any("nama kembar" in w for w in r["warnings"]))

    def test_build_ditolak_tanpa_confirm_dan_tidak_menimpa(self):
        self.assertNotEqual(run("build", "--input", self.d, "--keyword", "jun", "--out", self.out + "/x.xlsx").returncode, 0)
        os.makedirs(self.out); target = os.path.join(self.out, "x.xlsx"); open(target, "w").write("ada")
        self.assertNotEqual(run("build", "--input", self.d, "--keyword", "jun", "--out", target, "--confirm").returncode, 0)

    def test_build_verifikasi_rumus_dan_gaya(self):
        target = os.path.join(self.out, "hasil.xlsx")
        r = run("build", "--input", self.d, "--keyword", "jun", "--out", target, "--confirm")
        self.assertEqual(r.returncode, 0, r.stdout + r.stderr); self.assertIn("Status: LULUS", r.stdout)
        wb = openpyxl.load_workbook(target)
        self.assertEqual(wb.sheetnames, ["Index", "1 Jun", "1 Jun (1002)", "2 Jun", "3 Jun"])
        self.assertEqual(wb["1 Jun"]["A3"].value, "=SUM(A1:A2)")                 # rumus dalam sheet dipertahankan
        self.assertNotEqual(str(wb["1 Jun"]["A4"].value).startswith("="), True)   # rumus ke sheet lain jadi nilai
        self.assertFalse(wb["1 Jun"]["A1"].font.b); self.assertTrue(wb["1 Jun (1002)"]["A1"].font.b)   # gaya dari file berbeda tidak tertukar
        self.assertTrue(wb["1 Jun (1002)"]["A1"].font.color.rgb.endswith("0000FF"))
        self.assertEqual(wb["Index"]["B4"].hyperlink.location, "'1 Jun'!A1")
        self.assertEqual(list(wb["1 Jun"].merged_cells.ranges)[0].coord, "C1:D2")
        self.assertEqual(len(wb.defined_names), 0)

    def test_verify_mendeteksi_perubahan(self):
        target = os.path.join(self.out, "hasil.xlsx"); run("build", "--input", self.d, "--keyword", "jun", "--out", target, "--confirm")
        with zipfile.ZipFile(target) as z: parts = {n: z.read(n) for n in z.namelist()}
        parts["xl/worksheets/sheet2.xml"] = parts["xl/worksheets/sheet2.xml"].replace(b"<v>5</v>", b"<v>6</v>", 1)
        bad = os.path.join(self.out, "rusak.xlsx")
        with zipfile.ZipFile(bad, "w") as z: [z.writestr(n, b) for n, b in parts.items()]
        os.replace(target + ".report.json", bad + ".report.json")
        r = run("verify", "--input", self.d, "--keyword", "jun", "--out", bad)
        self.assertNotEqual(r.returncode, 0); self.assertIn("GAGAL", r.stdout)

    def test_is_cross(self):
        self.assertFalse(C.is_cross("SUM(A1:A2)", "1 Jun", set())); self.assertFalse(C.is_cross("'1 Jun'!A1+1", "1 Jun", set()))
        self.assertTrue(C.is_cross("'DPR '!P8", "1 Jun", set())); self.assertTrue(C.is_cross("[1]Sheet1!A1", "1 Jun", set()))
        self.assertTrue(C.is_cross("SUM(Total)", "1 Jun", {"total"})); self.assertTrue(C.is_cross("#REF!+1", "1 Jun", set()))


if __name__ == "__main__": unittest.main()
